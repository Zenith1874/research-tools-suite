#!/usr/bin/env python3
"""Read-only reconciliation of reviewed NFRA image cells and their source workbook."""
import argparse
import hashlib
import json
import math
import re
import sqlite3
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
ROW_LABELS = {
    'CN_BANK_NPL_BALANCE': '不良贷款余额',
    'CN_BANK_NPL_RATIO': '不良贷款率',
    'CN_BANK_NIM_YTD': '净息差',
    'CN_BANK_PROVISION_COVERAGE': '拨备覆盖率',
    'CN_BANK_CAPITAL_ADEQUACY': '资本充足率',
    'CN_BANK_TIER1_CAPITAL_ADEQUACY': '一级资本充足率',
    'CN_BANK_CORE_TIER1_CAPITAL_ADEQUACY': '核心一级资本充足率',
    'CN_BANK_NET_PROFIT_YTD': '净利润(本年累计)',
    'CN_BANK_ROA_YTD': '资产利润率',
    'CN_BANK_ROE_YTD': '资本利润率',
    'CN_BANK_LCR': '流动性覆盖率',
    'CN_BANK_LIQUIDITY_RATIO': '流动性比例',
    'CN_BANK_LOAN_DEPOSIT_RATIO': '存贷比(人民币)',
}


def normalized_label(value):
    return re.sub(r'\s+|\*', '', str(value or '')).replace('（', '(').replace('）', ')')


def reported_cell_value(value, number_format):
    """Match the source's displayed precision while retaining raw cells in the audit."""
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError('A finite numeric source cell is required; formulas are not evaluated.')
    number = Decimal(str(value))
    if number_format == '0.00%':
        return float((number * 100).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))
    if number_format.strip() in ('0', '0_'):
        return float(number.quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    raise ValueError(f'Unreviewed source number format: {number_format}')


def reviewed_workbook_cells(path, year):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook['商业银行季度']
        if f'{year}年' not in str(sheet['A1'].value):
            raise ValueError('Source heading does not confirm the requested statistical year.')
        if normalized_label(sheet['E2'].value) != '单位：亿元、%':
            raise ValueError('Source units do not match the reviewed contract.')
        if [sheet.cell(3, column).value for column in range(2, 6)] != ['一季度', '二季度', '三季度', '四季度']:
            raise ValueError('Source quarter columns have changed.')
        labels = {}
        for row in sheet.iter_rows(max_row=min(sheet.max_row, 250), max_col=1):
            label = normalized_label(row[0].value)
            if label in labels:
                raise ValueError('Duplicate source row label; reconciliation is ambiguous.')
            if label:
                labels[label] = row[0].row
        cells = []
        for code, label in ROW_LABELS.items():
            if label not in labels:
                raise ValueError(f'Required source row is missing: {label}')
            for quarter, column in enumerate(range(2, 6), 1):
                cell = sheet.cell(labels[label], column)
                reported = reported_cell_value(cell.value, cell.number_format)
                if reported is None:
                    raise ValueError(f'Required quarter cell is blank: {cell.coordinate}')
                cells.append({'indicator_code': code, 'period': f'{year}-{quarter * 3:02d}',
                              'sheet': sheet.title, 'cell': cell.coordinate,
                              'raw_value': cell.value, 'number_format': cell.number_format,
                              'reported_value': reported})
        return cells
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, required=True)
    parser.add_argument('--year', type=int, required=True)
    parser.add_argument('--history-path', type=Path, default=ROOT / 'data' / 'bank_supervision_history_reviewed.json')
    parser.add_argument('--db', type=Path, help='Optionally require an exact read-only database match.')
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'outputs' / 'macro_refresh_20261012')
    args = parser.parse_args()
    out = args.output_dir / ('bank-workbook-audit-' + datetime.now().strftime('%Y%m%d-%H%M%S-%f'))
    out.mkdir(parents=True, exist_ok=False)
    audit = {'captured_at': datetime.now().astimezone().isoformat(), 'statistical_year': args.year,
             'source_dir': str(args.source_dir), 'source_hash_checks': [], 'cells': [], 'errors': []}
    try:
        archive = json.loads((args.source_dir / 'audit.json').read_text(encoding='utf-8'))
        for source in archive['source_manifest']:
            path = (args.source_dir / source['raw_path']).resolve()
            if not path.is_relative_to(args.source_dir.resolve()):
                raise ValueError('Source archive path escaped its run directory.')
            actual = hashlib.sha256(path.read_bytes()).hexdigest()
            if actual != source['raw_sha256']:
                raise ValueError('Source archive checksum no longer matches its receipt.')
            audit['source_hash_checks'].append({'source_url': source['source_url'], 'sha256': actual})
        tables = json.loads(args.history_path.read_text(encoding='utf-8'))['tables']
        table = next(t for t in tables if t['year'] == args.year)
        source = next(s for s in archive['source_manifest'] if s['source_url'] == table['workbook_source_url'])
        if source['raw_sha256'] != table['workbook_sha256']:
            raise ValueError('Reviewed workbook vintage does not match the source archive.')
        audit['cells'] = reviewed_workbook_cells(args.source_dir / source['raw_path'], args.year)
        connection = sqlite3.connect(args.db.resolve().as_uri() + '?mode=ro', uri=True) if args.db else None
        try:
            for cell in audit['cells']:
                code, period = cell['indicator_code'], cell['period']
                quarter = int(period[-2:]) // 3 - 1
                cell['reviewed_value'] = table['values'][code][quarter]
                if abs(cell['reported_value'] - cell['reviewed_value']) > 1e-9:
                    raise ValueError(f'Workbook/image review mismatch: {code} {period}')
                if connection:
                    row = connection.execute('SELECT value,source_url FROM china_macro_observations WHERE '
                                             'indicator_code=? AND period=?', (code, period)).fetchone()
                    if not row or abs(row[0] - cell['reported_value']) > 1e-9 or row[1] != table['source_url']:
                        raise ValueError(f'Database/source mismatch: {code} {period}')
                    cell['database_value'] = row[0]
        finally:
            if connection:
                connection.close()
    except Exception as exc:
        audit['errors'].append(str(exc))
    audit['success'] = not audit['errors']
    (out / 'audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output_dir': str(out), 'success': audit['success'],
                      'cells_checked': len(audit['cells']), 'errors': audit['errors']}, ensure_ascii=False))
    raise SystemExit(0 if audit['success'] else 1)


if __name__ == '__main__':
    main()
