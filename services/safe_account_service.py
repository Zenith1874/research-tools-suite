"""Official SAFE BOP flows and IIP stocks with original units, signs and vintages."""
import io
import hashlib
import json
import math
import re
import unicodedata
from contextlib import closing
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin,urlsplit

PAGES={'bop':'https://www.safe.gov.cn/safe/2019/0627/13519.html',
       'iip':'https://www.safe.gov.cn/safe/2019/0627/13520.html'}

BOP_LABELS={
    'CA':'经常账户余额','GOODS_SERVICES':'货物和服务余额','GOODS':'货物贸易余额','SERVICES':'服务贸易余额',
    'PRIMARY_INCOME':'初次收入余额','SECONDARY_INCOME':'二次收入余额',
    'CAPITAL_FINANCIAL':'资本和金融账户余额','CAPITAL':'资本账户余额','FINANCIAL':'金融账户余额',
    'NONRESERVE':'非储备金融账户余额','DI':'直接投资净额','PORTFOLIO':'证券投资净额',
    'DERIVATIVES':'金融衍生工具净额','OTHER':'其他投资净额','RESERVES':'储备资产净流量',
    'ERRORS':'净误差与遗漏',
}
BOP_IDS={'1':'CA','1.A':'GOODS_SERVICES','1.A.a':'GOODS','1.A.b':'SERVICES',
         '1.B':'PRIMARY_INCOME','1.C':'SECONDARY_INCOME','2':'CAPITAL_FINANCIAL',
         '2.1':'CAPITAL','2.2':'FINANCIAL','2.2.1':'NONRESERVE','2.2.1.1':'DI',
         '2.2.1.2':'PORTFOLIO','2.2.1.3':'DERIVATIVES','2.2.1.4':'OTHER',
         '2.2.2':'RESERVES','3':'ERRORS'}
for parent in ('FINANCIAL','DI','PORTFOLIO','OTHER'):
    for side,label in (('ASSETS','资产'),('LIABILITIES','负债')):
        BOP_LABELS[parent+'_'+side]=BOP_LABELS[parent].replace('余额','').replace('净额','')+label+'净流量（官方符号）'

IIP_LABELS={'NET':'对外净头寸','ASSETS':'对外金融资产','LIABILITIES':'对外金融负债',
            'RESERVES_ASSETS':'储备资产头寸'}
for stem,label in (('DI','直接投资'),('PORTFOLIO','证券投资'),('DERIVATIVES','金融衍生工具'),('OTHER','其他投资')):
    for side,side_label in (('ASSETS','资产'),('LIABILITIES','负债')):
        IIP_LABELS[stem+'_'+side]=label+side_label+'头寸'
for sub,label in (('EQUITY','股权'),('DEBT','债务工具')):
    for side,side_label in (('ASSETS','资产'),('LIABILITIES','负债')):
        IIP_LABELS['DI_'+sub+'_'+side]='直接投资'+label+side_label+'头寸'

ACCOUNT_INDICATORS={}
for currency,unit in (('CNY','亿元'),('USD','亿美元')):
    for frequency,suffix in (('annual','A'),('quarterly','Q')):
        for stem,label in BOP_LABELS.items():
            ACCOUNT_INDICATORS[f'CN_SAFE_BOP_{stem}_{currency}_{suffix}']=(f'国际收支{label}（'+('年度' if suffix=='A' else '单季')+f'，{currency}）',unit,frequency)
for frequency,suffix in (('annual','A'),('quarterly','Q')):
    for stem,label in IIP_LABELS.items():
        ACCOUNT_INDICATORS[f'CN_SAFE_IIP_{stem}_USD_{suffix}']=(f'国际投资{label}（'+('年末' if suffix=='A' else '季末')+'，USD）','亿美元',frequency)
        ACCOUNT_INDICATORS[f'CN_SAFE_IIP_LEGACY_{stem}_USD_{suffix}']=(f'2015前旧口径国际投资{label}（'+('年末' if suffix=='A' else '季末')+'，USD）','亿美元',frequency)


def normalized(value):
    return re.sub(r'\s+','',unicodedata.normalize('NFKC',str(value or '')))


def period_from_header(value,frequency,kind):
    text=normalized(value)
    if frequency=='annual':
        match=re.fullmatch(r'((?:19|20)\d{2})(?:年末|年)?',text)
        return match[1] if match else None
    match=re.fullmatch(r'((?:19|20)\d{2})Q([1-4])',text)
    if match:return f'{match[1]}-{int(match[2])*3:02d}'
    match=re.fullmatch(r'((?:19|20)\d{2})年(?:(3|6|9)月)?末',text)
    return f'{match[1]}-{int(match[2] or 12):02d}' if match else None


def row_identity(label):
    text=normalized(label)
    match=re.match(r'^([0-9]+(?:\.[0-9A-Za-z]+)*)\.?(.*)$',text)
    return (match[1],match[2]) if match else ('',text)


def parse_safe_workbook(content,kind):
    """Return every numeric cell plus governed summary observations, without editing XLSX."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    if kind not in ('bop','iip'):raise ValueError('Unknown SAFE account kind.')
    book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    details=[];observations=[];sheet_audits=[]
    for sheet in book:
        frequency='annual' if '年度' in sheet.title else 'quarterly' if '季度' in sheet.title else None
        if frequency is None:continue
        currency='USD' if '美元' in sheet.title else 'CNY' if '人民币' in sheet.title else 'SDR' if 'SDR' in sheet.title else None
        if currency is None:continue
        rows=list(sheet.iter_rows(values_only=True))
        unit_text=''.join(str(cell or '') for row in rows[:5] for cell in row[:2])
        expected={'USD':'亿美元','CNY':'亿元人民币','SDR':'亿SDR'}[currency]
        if expected not in normalized(unit_text):raise ValueError(f'Unexpected unit in {sheet.title}.')
        header_index=next((i for i,row in enumerate(rows[:12]) if normalized(row[0])=='项目'),None)
        if header_index is None:raise ValueError(f'No explicit period header in {sheet.title}.')
        periods={j:period_from_header(value,frequency,kind) for j,value in enumerate(rows[header_index]) if j>0}
        periods={j:period for j,period in periods.items() if period}
        if not periods:
            book.close()
            raise ValueError(f'No explicit usable periods in {sheet.title}.')
        notes=[]
        in_notes=False
        for row in rows:
            text=' '.join(str(cell) for cell in row if cell is not None)
            if str(row[0] or '').startswith(('备注','说明')):in_notes=True
            if in_notes and text.strip():notes.append(text)
        parent_id='';section=''
        count=0
        for row_number,row in enumerate(rows[header_index+1:],header_index+2):
            if str(row[0] or '').startswith(('备注','说明')):break
            identifier,label=row_identity(row[0])
            if identifier:parent_id=identifier
            stem=None
            if kind=='bop':
                stem=BOP_IDS.get(identifier)
                for parent_id_key,parent_stem in (('2.2.1.1','DI'),('2.2.1.2','PORTFOLIO'),('2.2.1.4','OTHER')):
                    if identifier==parent_id_key+'.1' and label=='资产':stem=parent_stem+'_ASSETS'
                    if identifier==parent_id_key+'.2' and label=='负债':stem=parent_stem+'_LIABILITIES'
                if not identifier and label in ('资产','负债'):
                    parent=BOP_IDS.get(parent_id)
                    if parent in ('FINANCIAL','DI','PORTFOLIO','OTHER'):
                        stem=parent+('_ASSETS' if label=='资产' else '_LIABILITIES')
            else:
                if not identifier and label in ('资产','负债'):
                    section='ASSETS' if label=='资产' else 'LIABILITIES';stem=section
                elif not identifier and label=='净头寸':stem='NET'
                elif identifier in ('1','2','3','4') and section:
                    stem={'1':'DI','2':'PORTFOLIO','3':'DERIVATIVES','4':'OTHER'}[identifier]+'_'+section
                elif identifier in ('1.1','1.2') and section:
                    stem='DI_'+('EQUITY' if identifier=='1.1' else 'DEBT')+'_'+section
                elif identifier=='5' and section=='ASSETS':stem='RESERVES_ASSETS'
            for column,period in periods.items():
                if column>=len(row):continue
                value=row[column]
                if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):continue
                item={'sheet':sheet.title,'cell':f'{get_column_letter(column+1)}{row_number}',
                      'row_id':identifier or parent_id,'row_label':str(row[0] or ''),'period':period,
                      'currency':currency,'unit':expected,'frequency':frequency,'value':float(value)}
                details.append(item);count+=1
                if stem and currency in ('USD','CNY'):
                    suffix='A' if frequency=='annual' else 'Q'
                    legacy='LEGACY_' if kind=='iip' and int(period[:4])<2015 else ''
                    code=f'CN_SAFE_{kind.upper()}_{legacy}{stem}_{currency}_{suffix}'
                    if code in ACCOUNT_INDICATORS:observations.append({**item,'indicator_code':code})
        sheet_audits.append({'sheet':sheet.title,'frequency':frequency,'currency':currency,
                            'first':min(periods.values()),'last':max(periods.values()),
                            'numeric_cells':count,'notes':notes})
    book.close()
    keys=[(row['indicator_code'],row['period']) for row in observations]
    if len(keys)!=len(set(keys)):raise ValueError('Duplicate governed summary keys in source workbook.')
    return {'details':details,'observations':observations,'sheets':sheet_audits}


def ensure_safe_account_tables(conn):
    conn.execute('''CREATE TABLE IF NOT EXISTS safe_account_cells (
        raw_sha256 TEXT,source_kind TEXT,sheet_name TEXT,cell_ref TEXT,row_id TEXT,row_label TEXT,
        period TEXT,frequency TEXT,currency TEXT,unit TEXT,value REAL,source_url TEXT,release_date TEXT,
        retrieved_at TEXT,UNIQUE(raw_sha256,sheet_name,cell_ref))''')
    conn.execute('''CREATE TABLE IF NOT EXISTS safe_account_revisions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,indicator_code TEXT,period TEXT,old_row_json TEXT,
        new_value REAL,new_source_url TEXT,changed_at TEXT)''')
    conn.execute('''CREATE TABLE IF NOT EXISTS safe_account_versions (
        raw_sha256 TEXT PRIMARY KEY,source_kind TEXT,source_url TEXT,release_date TEXT,
        retrieved_at TEXT,sheet_audit_json TEXT)''')
    conn.commit()


def update_safe_accounts(db_path,source_dir=None,session=None):
    """Use the official revision time series, preserving every raw vintage and prior row."""
    from bs4 import BeautifulSoup
    from services.china_macro_service import connect,ensure_china_macro_tables
    import requests
    started=datetime.now().astimezone().isoformat()
    if source_dir:
        directory=Path(source_dir)
        manifest=json.loads((directory/'audit.json').read_text(encoding='utf-8'))['source_manifest']
    else:
        directory=Path(__file__).resolve().parents[1]/'data'/'china_macro_raw'/('safe_accounts_'+datetime.now().strftime('%Y%m%d-%H%M%S-%f'))
        directory.mkdir(parents=True,exist_ok=False)
        manifest=[]
    session=session or requests.Session()
    changed=0;raw_inserted=0;errors=[];reports=[]
    def read_or_fetch(url,name):
        path=directory/name
        if source_dir:
            prior=next(item for item in manifest if item['raw_path']==name)
            content=path.read_bytes()
            if hashlib.sha256(content).hexdigest()!=prior['raw_sha256']:
                raise ValueError('Archived official account file hash mismatch.')
            return content,prior
        response=session.get(url,timeout=20,headers={'User-Agent':'Mozilla/5.0'})
        response.raise_for_status()
        if urlsplit(response.url).hostname!='www.safe.gov.cn':raise ValueError('Non-SAFE source redirect.')
        content=response.content;path.write_bytes(content)
        item={'source_url':url,'raw_path':name,'raw_sha256':hashlib.sha256(content).hexdigest(),
              'retrieved_at':datetime.now().astimezone().isoformat()}
        manifest.append(item)
        with (directory/'manifest.jsonl').open('a',encoding='utf-8') as handle:
            handle.write(json.dumps(item,ensure_ascii=False)+'\n')
        return content,item
    with closing(connect(db_path)) as conn:
        ensure_china_macro_tables(conn);ensure_safe_account_tables(conn)
        for kind,page in PAGES.items():
            try:
                page_bytes,page_record=read_or_fetch(page,f'{kind}-page.html')
                soup=BeautifulSoup(page_bytes.decode('utf-8'),'html.parser')
                match=re.search(r'发布日期[：:]?(20\d{2}-\d{2}-\d{2})',normalized(soup.get_text('',strip=True)))
                release_date=match[1] if match else None
                links=[urljoin(page,a['href']) for a in soup.find_all('a',href=True)
                       if urlsplit(urljoin(page,a['href'])).path.endswith('.xlsx')]
                if len(set(links))!=1:raise ValueError('Ambiguous official workbook link.')
                content,record=read_or_fetch(links[0],f'{kind}.xlsx')
                parsed=parse_safe_workbook(content,kind)
                now=record['retrieved_at'];digest=record['raw_sha256']
                raw_rows=[(digest,kind,row['sheet'],row['cell'],row['row_id'],row['row_label'],row['period'],
                           row['frequency'],row['currency'],row['unit'],row['value'],links[0],release_date,now)
                          for row in parsed['details']]
                before=conn.total_changes
                conn.executemany('INSERT OR IGNORE INTO safe_account_cells VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',raw_rows)
                book_raw_inserted=conn.total_changes-before
                book_changed=0
                for row in parsed['observations']:
                    code=row['indicator_code'];period=row['period'];value=row['value']
                    old=conn.execute('SELECT * FROM china_macro_observations WHERE indicator_code=? AND period=?',(code,period)).fetchone()
                    if old and old['release_date'] and release_date and old['release_date']>release_date:
                        continue
                    if old and old['value']==value and old['source_url']==links[0] and old['release_date']==release_date:
                        continue
                    if old:
                        conn.execute('INSERT INTO safe_account_revisions (indicator_code,period,old_row_json,new_value,new_source_url,changed_at) VALUES (?,?,?,?,?,?)',
                                     (code,period,json.dumps(dict(old),ensure_ascii=False),value,links[0],now))
                    name,unit,frequency=ACCOUNT_INDICATORS[code]
                    notes=('SAFE official revision workbook; original currency columns and accounting signs preserved. '
                           'BOP assets net increase is negative; liabilities net increase is positive. Not MOFCOM utilized FDI. '
                           'IIP is a stock, not a flow; pre-2015 IIP is a separate legacy series; trade-credit survey break at 2017 year-end.'
                           f' Source sheet={row["sheet"]}; cell={row["cell"]}; SHA256={digest}.')
                    conn.execute('''INSERT INTO china_macro_observations
                        (indicator_code,indicator_name,period,value,unit,frequency,data_status,source_name,source_type,
                         source_url,source_title,parser_notes,updated_at,release_date,retrieved_at,parser_version)
                        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                        ON CONFLICT(indicator_code,period) DO UPDATE SET value=excluded.value,
                        data_status=excluded.data_status,source_url=excluded.source_url,source_title=excluded.source_title,
                        parser_notes=excluded.parser_notes,updated_at=excluded.updated_at,
                        release_date=excluded.release_date,retrieved_at=excluded.retrieved_at,parser_version=excluded.parser_version''',
                        (code,name,period,value,unit,frequency,'official','国家外汇管理局',f'safe_{kind}_revision_workbook',
                         links[0],f'SAFE {kind.upper()} / {row["sheet"]}',notes,now,release_date,now,'safe-accounts-20261010'))
                    book_changed+=1
                conn.execute('INSERT OR IGNORE INTO safe_account_versions VALUES (?,?,?,?,?,?)',
                             (digest,kind,links[0],release_date,now,json.dumps(parsed['sheets'],ensure_ascii=False)))
                conn.commit();changed+=book_changed;raw_inserted+=book_raw_inserted
                reports.append({'kind':kind,'observations':len(parsed['observations']),'changed':book_changed,
                                'raw_numeric_cells':len(parsed['details']),'sheets':parsed['sheets']})
            except Exception as exc:
                conn.rollback();errors.append({'kind':kind,'error':str(exc)})
    result={'started_at':started,'finished_at':datetime.now().astimezone().isoformat(),
            'success':not errors,'records_upserted':changed,'raw_cells_inserted':raw_inserted,
            'reports':reports,'errors':errors,'raw_directory':str(directory)}
    if not source_dir:
        (directory/'audit.json').write_text(json.dumps({'source_manifest':manifest,'result':result},ensure_ascii=False,indent=2),encoding='utf-8')
    return result
