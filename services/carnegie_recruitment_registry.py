# -*- coding: utf-8 -*-
"""Carnegie 2025 recruitment-portal registry.

This registry is deliberately separate from the older IPEDS 2023
R1/R2/Doctoral-Professional registry.  Its target grain is one Carnegie
institution (UNITID), while public-site requests are de-duplicated by the
normalized official website host.

The default scope is:

* every 2025 Research 2 institution; and
* every four-year institution whose 2025 Institutional Classification is a
  Baccalaureate or Undergraduate/Graduate-Master's category.

The second group is an explicitly documented *teaching-oriented proxy*, not
an official Carnegie "teaching school" designation.  Special Focus: Business
institutions are counted and retained as an optional extension, but are not
part of the default first batch.

The site scanner is shared with the older registry.  It discovers public
recruitment entry points only; a portal candidate is not vacancy evidence.
"""

from __future__ import annotations

import csv
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import time
import urllib.request

from services.ipeds_recruitment_registry import (
    RecruitmentSiteScanner,
    SiteScanResult,
    _exception_status,
    _json,
    _short,
    canonical_host,
    normalize_public_url,
    now_utc,
    reset_crawl_budgets,
    reset_robots_cache,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE_DB = ROOT / "data" / "academic_market.db"
DEFAULT_CARNEGIE_DIR = ROOT / "data" / "academic_market" / "releases" / "carnegie"
DEFAULT_PUBLIC_FILE = DEFAULT_CARNEGIE_DIR / "2025-Public-Data-File.xlsx"
DEFAULT_RAD_FILE = DEFAULT_CARNEGIE_DIR / "2025-RAD-Public-Data-File.xlsx"
DEFAULT_REGISTRY_DB = ROOT / "data" / "faculty_jobs" / "carnegie_2025_recruitment_registry.db"
DEFAULT_EXPORT_CSV = ROOT / "outputs" / "carnegie_2025_recruitment_registry_20260906.csv"

CARNEGIE_PUBLIC_URL = (
    "https://carnegieclassifications.acenet.edu/wp-content/uploads/2025/04/"
    "2025-Public-Data-File.xlsx"
)
CARNEGIE_RAD_URL = (
    "https://carnegieclassifications.acenet.edu/wp-content/uploads/2025/02/"
    "2025-RAD-Public-Data-File.xlsx"
)
IPEDS_WEBSITE_YEAR = 2023
CARNEGIE_DATA_YEAR = 2025
REGISTRY_SCHEMA_VERSION = 1
DISCOVERY_RULE_VERSION = "carnegie_2025_r2_teaching_v1_20260906"
DEFAULT_SEED_SCOPE = "CARNEGIE_2025_R2_PLUS_TEACHING_ORIENTED_4YR"
BUSINESS_SCOPE_SUFFIX = "_PLUS_BUSINESS_FOCUSED"
USER_AGENT = (
    "research-tools-suite/carnegie-registry (+local research workstation; "
    "public Carnegie data download)"
)

R2_NAME = "Research 2: High Spending and Doctorate Production"
RCU_NAME = "Research Colleges and Universities"
R1_NAME = "Research 1: Very High Spending and Doctorate Production"
BUSINESS_FOCUS_NAME = "Special Focus: Business"


def _int(value):
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _text(value):
    value = str(value or "").strip()
    return value or None


def _get(row, key, default=None):
    if row is None:
        return default
    if hasattr(row, "keys"):
        try:
            if key in row.keys():
                return row[key]
        except (AttributeError, TypeError):
            pass
    if isinstance(row, dict):
        return row.get(key, default)
    return default


def is_r2(research_name):
    """Return whether a Carnegie 2025 research designation is R2."""
    return _text(research_name) == R2_NAME


def is_teaching_oriented_classification(classification_name):
    """Apply the documented four-year teaching-oriented proxy.

    Associate/Baccalaureate categories are excluded because the recruitment
    target is a university/college four-year faculty market, not the broader
    community-college and mixed associate market.
    """
    name = _text(classification_name) or ""
    if "Associate/Baccalaureate" in name:
        return False
    return (
        "Baccalaureate" in name
        or "Undergraduate/Graduate-Master's" in name
    )


def _target_segment(*, r2, teaching, business):
    if r2 and teaching:
        return "r2_and_teaching_oriented"
    if r2:
        return "r2"
    if teaching:
        return "teaching_oriented"
    if business:
        return "business_focused_special_focus"
    return None


def _read_xlsx_rows(path, sheet_name):
    """Read one public Carnegie sheet as dictionaries without changing it."""
    try:
        import openpyxl
    except ImportError as error:  # pragma: no cover - environment-specific
        raise RuntimeError(
            "openpyxl is required to read the official Carnegie XLSX files."
        ) from error

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"{path} does not contain sheet {sheet_name!r}")
        rows = workbook[sheet_name].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError(f"{path} sheet {sheet_name!r} has no header")
        names = [str(item).strip() if item is not None else "" for item in header]
        if not names[0]:
            raise ValueError(f"{path} sheet {sheet_name!r} has an empty first column")
        result = []
        for values in rows:
            if not values or all(value is None for value in values):
                continue
            result.append({name: value for name, value in zip(names, values) if name})
        return result
    finally:
        workbook.close()


def _file_metadata(path, source_url, source_key):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(
            f"Missing Carnegie source file: {path}. Download it from {source_url} "
            "or pass --download to the seed script."
        )
    hasher = hashlib.sha256()
    size = 0
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            size += len(chunk)
            hasher.update(chunk)
    return {
        "source_key": source_key,
        "source_url": source_url,
        "local_path": str(path),
        "downloaded_at": now_utc(),
        "byte_size": size,
        "sha256": hasher.hexdigest(),
        "http_status": 200,
        "status": "downloaded",
    }


def _download(url, destination, *, overwrite=False, timeout=120):
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and not overwrite:
        return destination
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=timeout) as response:
        payload = response.read()
        status = getattr(response, "status", 200)
    if not payload:
        raise IOError(f"{url} returned {status} with an empty body")
    destination.write_bytes(payload)
    return destination


def download_carnegie_sources(*, public_path=DEFAULT_PUBLIC_FILE,
                              rad_path=DEFAULT_RAD_FILE, overwrite=False):
    """Download the two official 2025 XLSX files, preserving raw bytes."""
    public = _download(CARNEGIE_PUBLIC_URL, public_path, overwrite=overwrite)
    rad = _download(CARNEGIE_RAD_URL, rad_path, overwrite=overwrite)
    return {
        "public": _file_metadata(public, CARNEGIE_PUBLIC_URL, "carnegie_2025_public"),
        "rad": _file_metadata(rad, CARNEGIE_RAD_URL, "carnegie_2025_rad"),
    }


def _scope_name(include_business_focused):
    return DEFAULT_SEED_SCOPE + (BUSINESS_SCOPE_SUFFIX if include_business_focused else "")


def build_target_rows(public_rows, rad_rows, ipeds_rows, *,
                      include_business_focused=False, seed_scope=None):
    """Build target rows and diagnostics from official rows plus IPEDS URLs.

    ``public_rows`` is the 2025 Public Data File's ``data`` sheet,
    ``rad_rows`` is the RAD file's ``Data`` sheet, and ``ipeds_rows`` is a
    mapping or iterable from the local IPEDS HD table.  No rows are written by
    this function, which makes the selection rule independently testable.
    """
    scope = seed_scope or _scope_name(include_business_focused)
    rad_by_unitid = {
        _int(_get(row, "UNITID")): row for row in rad_rows
        if _int(_get(row, "UNITID")) is not None
    }
    if isinstance(ipeds_rows, dict):
        ipeds_by_unitid = ipeds_rows
    else:
        ipeds_by_unitid = {
            _int(_get(row, "unitid")): row for row in ipeds_rows
            if _int(_get(row, "unitid")) is not None
        }

    all_r2 = set()
    all_teaching = set()
    all_business = set()
    selected = []
    for public in public_rows:
        unitid = _int(_get(public, "unitid"))
        if unitid is None:
            continue
        rad = rad_by_unitid.get(unitid)
        research_name = (
            _text(_get(public, "research2025name"))
            or _text(_get(rad, "2025 Research Activity Designation"))
        )
        classification = _text(_get(public, "ic2025name"))
        r2 = is_r2(research_name)
        teaching = is_teaching_oriented_classification(classification)
        business = classification == BUSINESS_FOCUS_NAME
        if r2:
            all_r2.add(unitid)
        if teaching:
            all_teaching.add(unitid)
        if business:
            all_business.add(unitid)
        if not (r2 or teaching or (include_business_focused and business)):
            continue

        ipeds = ipeds_by_unitid.get(unitid)
        raw_website = _text(_get(ipeds, "website"))
        official_url = normalize_public_url(raw_website)
        host = canonical_host(official_url) if official_url else None
        local_state = _text(_get(ipeds, "state"))
        state = local_state or _text(_get(public, "stabbr")) or "UNKNOWN"
        segment = _target_segment(r2=r2, teaching=teaching, business=business)
        selected.append({
            "unitid": unitid,
            "data_year": CARNEGIE_DATA_YEAR,
            "institution_name": _text(_get(public, "instnm")) or str(unitid),
            "city": _text(_get(public, "city")),
            "state": state,
            "control": _int(_get(public, "control"))
                       if _get(public, "control") is not None
                       else _int(_get(ipeds, "control")),
            "carnegie_research_code": _int(_get(public, "research2025")),
            "carnegie_research_designation": research_name or "No research designation",
            "carnegie_institutional_code": _int(_get(public, "ic2025")),
            "carnegie_institutional_classification": classification or "Unknown",
            "is_r2": int(r2),
            "is_teaching_oriented": int(teaching),
            "teaching_proxy_rule": (
                "2025 IC Baccalaureate or Undergraduate/Graduate-Master's; "
                "Associate/Baccalaureate excluded"
            ) if teaching else None,
            "is_business_focused": int(business),
            "target_segment": segment,
            "official_website_raw": raw_website,
            "official_website_url": official_url,
            "official_host": host,
            "website_group_key": host,
            "seed_scope": scope,
            "seed_source": (
                "Carnegie 2025 Public Data File + 2025 RAD Public Data File + "
                f"IPEDS HD {IPEDS_WEBSITE_YEAR} website"
            ),
            "source_release_id": _int(_get(ipeds, "release_id")),
        })

    diagnostics = {
        "public_rows": len({
            _int(_get(row, "unitid")) for row in public_rows
            if _int(_get(row, "unitid")) is not None
        }),
        "rad_rows": len(rad_by_unitid),
        "r2_count": len(all_r2),
        "teaching_oriented_count": len(all_teaching),
        "business_focused_available": len(all_business),
        "r2_teaching_overlap": len(all_r2 & all_teaching),
        "target_count": len(selected),
        "scope": scope,
        "include_business_focused": bool(include_business_focused),
    }
    return selected, diagnostics


def registry_connect(db_path=None):
    path = Path(db_path or DEFAULT_REGISTRY_DB)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(path), timeout=60)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def ensure_registry_schema(conn):
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS registry_migrations (
            version INTEGER PRIMARY KEY,
            applied_at TEXT NOT NULL,
            note TEXT
        );

        CREATE TABLE IF NOT EXISTS registry_sources (
            source_key TEXT PRIMARY KEY,
            source_url TEXT NOT NULL,
            local_path TEXT NOT NULL,
            downloaded_at TEXT NOT NULL,
            byte_size INTEGER,
            sha256 TEXT,
            http_status INTEGER,
            status TEXT NOT NULL,
            note TEXT
        );

        CREATE TABLE IF NOT EXISTS registry_runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            mode TEXT NOT NULL,
            data_year INTEGER,
            seed_scope TEXT,
            started_at TEXT NOT NULL,
            finished_at TEXT,
            status TEXT NOT NULL,
            requested_count INTEGER NOT NULL DEFAULT 0,
            processed_count INTEGER NOT NULL DEFAULT 0,
            request_attempts INTEGER NOT NULL DEFAULT 0,
            result_counts_json TEXT NOT NULL DEFAULT '{}',
            source_db TEXT,
            rule_version TEXT,
            error TEXT
        );

        CREATE TABLE IF NOT EXISTS institution_recruitment_websites (
            website_group_key TEXT PRIMARY KEY,
            representative_unitid INTEGER NOT NULL,
            official_website_url TEXT,
            official_host TEXT,
            linked_institution_count INTEGER NOT NULL DEFAULT 0,
            first_seeded_at TEXT NOT NULL,
            last_seeded_at TEXT NOT NULL,
            latest_status TEXT NOT NULL DEFAULT 'seeded',
            latest_checked_at TEXT,
            latest_http_status INTEGER,
            latest_portals_found INTEGER NOT NULL DEFAULT 0,
            latest_reason TEXT
        );

        CREATE TABLE IF NOT EXISTS institution_recruitment_targets (
            unitid INTEGER PRIMARY KEY,
            data_year INTEGER NOT NULL,
            institution_name TEXT NOT NULL,
            city TEXT,
            state TEXT NOT NULL,
            control INTEGER,
            carnegie_research_code INTEGER,
            carnegie_research_designation TEXT NOT NULL,
            carnegie_institutional_code INTEGER,
            carnegie_institutional_classification TEXT NOT NULL,
            is_r2 INTEGER NOT NULL DEFAULT 0,
            is_teaching_oriented INTEGER NOT NULL DEFAULT 0,
            teaching_proxy_rule TEXT,
            is_business_focused INTEGER NOT NULL DEFAULT 0,
            target_segment TEXT NOT NULL,
            official_website_raw TEXT,
            official_website_url TEXT,
            official_host TEXT,
            website_group_key TEXT,
            seed_scope TEXT NOT NULL,
            seed_source TEXT NOT NULL,
            source_release_id INTEGER,
            seeded_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (website_group_key)
                REFERENCES institution_recruitment_websites(website_group_key)
        );
        CREATE INDEX IF NOT EXISTS idx_carnegie_targets_segment
            ON institution_recruitment_targets(target_segment);
        CREATE INDEX IF NOT EXISTS idx_carnegie_targets_website
            ON institution_recruitment_targets(website_group_key);

        CREATE TABLE IF NOT EXISTS institution_recruitment_portals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            website_group_key TEXT NOT NULL,
            portal_url TEXT NOT NULL,
            canonical_url TEXT NOT NULL,
            origin_host TEXT,
            portal_kind TEXT NOT NULL,
            discovery_method TEXT NOT NULL,
            discovery_evidence TEXT,
            confidence TEXT NOT NULL,
            first_seen_at TEXT NOT NULL,
            last_seen_at TEXT NOT NULL,
            latest_page_status TEXT NOT NULL DEFAULT 'not_checked',
            latest_http_status INTEGER,
            latest_final_url TEXT,
            latest_title TEXT,
            latest_jobposting_jsonld INTEGER NOT NULL DEFAULT 0,
            active INTEGER NOT NULL DEFAULT 1,
            UNIQUE (website_group_key, canonical_url),
            FOREIGN KEY (website_group_key)
                REFERENCES institution_recruitment_websites(website_group_key)
        );
        CREATE INDEX IF NOT EXISTS idx_carnegie_portals_group
            ON institution_recruitment_portals(website_group_key, active);

        CREATE TABLE IF NOT EXISTS institution_recruitment_observations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id INTEGER,
            website_group_key TEXT NOT NULL,
            checked_at TEXT NOT NULL,
            official_website_url TEXT,
            homepage_status TEXT NOT NULL,
            homepage_http_status INTEGER,
            homepage_final_url TEXT,
            overall_status TEXT NOT NULL,
            candidates_seen INTEGER NOT NULL DEFAULT 0,
            portals_found INTEGER NOT NULL DEFAULT 0,
            pages_fetched INTEGER NOT NULL DEFAULT 0,
            request_attempts INTEGER NOT NULL DEFAULT 0,
            discovery_complete INTEGER NOT NULL DEFAULT 0,
            candidate_truncated INTEGER NOT NULL DEFAULT 0,
            reason TEXT,
            evidence_json TEXT NOT NULL DEFAULT '{}',
            rule_version TEXT NOT NULL,
            FOREIGN KEY (run_id) REFERENCES registry_runs(id),
            FOREIGN KEY (website_group_key)
                REFERENCES institution_recruitment_websites(website_group_key)
        );
        CREATE INDEX IF NOT EXISTS idx_carnegie_observations_group
            ON institution_recruitment_observations(website_group_key, checked_at);
        CREATE INDEX IF NOT EXISTS idx_carnegie_observations_status
            ON institution_recruitment_observations(overall_status, checked_at);
        """
    )
    conn.execute(
        "INSERT OR IGNORE INTO registry_migrations(version, applied_at, note) VALUES (?,?,?)",
        (
            REGISTRY_SCHEMA_VERSION,
            now_utc(),
            "Carnegie 2025 R2 plus documented four-year teaching-oriented proxy "
            "recruitment portal registry; separate from the legacy IPEDS registry.",
        ),
    )
    conn.commit()


def _source_connect(path):
    source = Path(path)
    conn = sqlite3.connect(f"file:{source.as_posix()}?mode=ro", uri=True, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def _upsert_source(conn, metadata):
    conn.execute(
        """INSERT INTO registry_sources
           (source_key,source_url,local_path,downloaded_at,byte_size,sha256,
            http_status,status,note)
           VALUES (?,?,?,?,?,?,?,?,?)
           ON CONFLICT(source_key) DO UPDATE SET
             source_url=excluded.source_url,local_path=excluded.local_path,
             downloaded_at=excluded.downloaded_at,byte_size=excluded.byte_size,
             sha256=excluded.sha256,http_status=excluded.http_status,
             status=excluded.status,note=excluded.note""",
        (
            metadata["source_key"], metadata["source_url"], metadata["local_path"],
            metadata["downloaded_at"], metadata["byte_size"], metadata["sha256"],
            metadata["http_status"], metadata["status"],
            "Official Carnegie 2025 public data file; raw bytes retained locally.",
        ),
    )


def _current_scope(conn):
    row = conn.execute(
        "SELECT seed_scope FROM institution_recruitment_targets "
        "ORDER BY updated_at DESC LIMIT 1"
    ).fetchone()
    return row[0] if row else DEFAULT_SEED_SCOPE


def seed_registry(source_db=None, registry_db=None, *,
                  public_file=DEFAULT_PUBLIC_FILE, rad_file=DEFAULT_RAD_FILE,
                  include_business_focused=False, validate_release=True):
    """Seed the current Carnegie 2025 R2/teaching-oriented target registry."""
    source_path = Path(source_db or DEFAULT_SOURCE_DB)
    registry_path = Path(registry_db or DEFAULT_REGISTRY_DB)
    public_path = Path(public_file)
    rad_path = Path(rad_file)
    scope = _scope_name(include_business_focused)
    public_meta = _file_metadata(public_path, CARNEGIE_PUBLIC_URL, "carnegie_2025_public")
    rad_meta = _file_metadata(rad_path, CARNEGIE_RAD_URL, "carnegie_2025_rad")

    public_rows = _read_xlsx_rows(public_path, "data")
    rad_rows = _read_xlsx_rows(rad_path, "Data")
    public_unitids = {
        _int(_get(row, "unitid")) for row in public_rows
        if _int(_get(row, "unitid")) is not None
    }
    rad_unitids = {
        _int(_get(row, "UNITID")) for row in rad_rows
        if _int(_get(row, "UNITID")) is not None
    }
    if validate_release:
        if len(public_unitids) != 3927:
            raise ValueError(
                f"Unexpected Carnegie 2025 public institution count: {len(public_unitids)}; "
                "refusing to seed a possibly partial file."
            )
        if len(rad_unitids) != 542:
            raise ValueError(
                f"Unexpected Carnegie 2025 RAD institution count: {len(rad_unitids)}; "
                "refusing to seed a possibly partial file."
            )

    source = _source_connect(source_path)
    conn = registry_connect(registry_path)
    ensure_registry_schema(conn)
    started = now_utc()
    run_id = conn.execute(
        """INSERT INTO registry_runs
           (mode,data_year,seed_scope,started_at,status,source_db,rule_version)
           VALUES (?,?,?,?,?,?,?)""",
        ("seed", CARNEGIE_DATA_YEAR, scope, started, "running",
         str(source_path), DISCOVERY_RULE_VERSION),
    ).lastrowid
    try:
        ipeds_rows = source.execute(
            """SELECT unitid,institution_name,state,control,website,release_id
               FROM ipeds_institutions WHERE data_year=?""",
            (IPEDS_WEBSITE_YEAR,),
        ).fetchall()
        selected, diagnostics = build_target_rows(
            public_rows, rad_rows, ipeds_rows,
            include_business_focused=include_business_focused,
            seed_scope=scope,
        )
        if validate_release and diagnostics["r2_count"] != 139:
            raise ValueError(
                f"Unexpected Carnegie 2025 R2 count: {diagnostics['r2_count']}; "
                "refusing to seed a mismatched release."
            )

        with conn:
            _upsert_source(conn, public_meta)
            _upsert_source(conn, rad_meta)
            seeded_at = now_utc()
            representatives = {}
            for row in selected:
                group_key = row["website_group_key"]
                if not group_key:
                    continue
                rank = (
                    0 if row["is_r2"] else 1,
                    row["institution_name"].casefold(),
                    row["unitid"],
                )
                previous = representatives.get(group_key)
                if previous is None or rank < previous[0]:
                    representatives[group_key] = (rank, row)

            for group_key, (_, row) in representatives.items():
                conn.execute(
                    """INSERT INTO institution_recruitment_websites
                       (website_group_key,representative_unitid,official_website_url,
                        official_host,linked_institution_count,first_seeded_at,last_seeded_at)
                       VALUES (?,?,?,?,0,?,?)
                       ON CONFLICT(website_group_key) DO UPDATE SET
                         representative_unitid=excluded.representative_unitid,
                         official_website_url=excluded.official_website_url,
                         official_host=excluded.official_host,
                         last_seeded_at=excluded.last_seeded_at""",
                    (group_key, row["unitid"], row["official_website_url"],
                     row["official_host"], seeded_at, seeded_at),
                )

            for row in selected:
                conn.execute(
                    """INSERT INTO institution_recruitment_targets
                       (unitid,data_year,institution_name,city,state,control,
                        carnegie_research_code,carnegie_research_designation,
                        carnegie_institutional_code,carnegie_institutional_classification,
                        is_r2,is_teaching_oriented,teaching_proxy_rule,is_business_focused,
                        target_segment,official_website_raw,official_website_url,
                        official_host,website_group_key,seed_scope,seed_source,
                        source_release_id,seeded_at,updated_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                       ON CONFLICT(unitid) DO UPDATE SET
                         data_year=excluded.data_year,institution_name=excluded.institution_name,
                         city=excluded.city,state=excluded.state,control=excluded.control,
                         carnegie_research_code=excluded.carnegie_research_code,
                         carnegie_research_designation=excluded.carnegie_research_designation,
                         carnegie_institutional_code=excluded.carnegie_institutional_code,
                         carnegie_institutional_classification=excluded.carnegie_institutional_classification,
                         is_r2=excluded.is_r2,is_teaching_oriented=excluded.is_teaching_oriented,
                         teaching_proxy_rule=excluded.teaching_proxy_rule,
                         is_business_focused=excluded.is_business_focused,
                         target_segment=excluded.target_segment,
                         official_website_raw=excluded.official_website_raw,
                         official_website_url=excluded.official_website_url,
                         official_host=excluded.official_host,website_group_key=excluded.website_group_key,
                         seed_scope=excluded.seed_scope,seed_source=excluded.seed_source,
                         source_release_id=excluded.source_release_id,updated_at=excluded.updated_at""",
                    (
                        row["unitid"], row["data_year"], row["institution_name"],
                        row["city"], row["state"], row["control"],
                        row["carnegie_research_code"], row["carnegie_research_designation"],
                        row["carnegie_institutional_code"],
                        row["carnegie_institutional_classification"], row["is_r2"],
                        row["is_teaching_oriented"], row["teaching_proxy_rule"],
                        row["is_business_focused"], row["target_segment"],
                        row["official_website_raw"], row["official_website_url"],
                        row["official_host"], row["website_group_key"], row["seed_scope"],
                        row["seed_source"], row["source_release_id"], seeded_at, seeded_at,
                    ),
                )
            conn.execute(
                """UPDATE institution_recruitment_websites AS w
                   SET linked_institution_count=(
                     SELECT COUNT(*) FROM institution_recruitment_targets AS t
                     WHERE t.website_group_key=w.website_group_key
                   )"""
            )
            website_count = conn.execute(
                """SELECT COUNT(*) FROM institution_recruitment_websites
                   WHERE linked_institution_count>0"""
            ).fetchone()[0]
            missing_website_count = sum(
                row["official_website_url"] is None for row in selected
            )
            result = {
                **diagnostics,
                "website_group_count": website_count,
                "missing_official_website": missing_website_count,
                "duplicate_website_records": max(
                    0, len(selected) - website_count - missing_website_count
                ),
                "source_files": {
                    "public": public_meta,
                    "rad": rad_meta,
                },
                "registry_db": str(registry_path),
            }
            conn.execute(
                """UPDATE registry_runs SET finished_at=?,status='ok',requested_count=?,
                   processed_count=?,result_counts_json=? WHERE id=?""",
                (now_utc(), len(selected), len(selected), _json(result), run_id),
            )
        return result
    except Exception as error:
        conn.execute(
            """UPDATE registry_runs SET finished_at=?,status='failed',error=? WHERE id=?""",
            (now_utc(), _short(error, 1000), run_id),
        )
        conn.commit()
        raise
    finally:
        source.close()
        conn.close()


def _record_site_scan(conn, result: SiteScanResult, *, run_id):
    seen_urls = {portal["canonical_url"] for portal in result.portals}
    for portal in result.portals:
        now = result.checked_at
        conn.execute(
            """INSERT INTO institution_recruitment_portals
               (website_group_key,portal_url,canonical_url,origin_host,portal_kind,
                discovery_method,discovery_evidence,confidence,first_seen_at,last_seen_at,
                latest_page_status,latest_http_status,latest_final_url,latest_title,
                latest_jobposting_jsonld)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(website_group_key,canonical_url) DO UPDATE SET
                 portal_url=excluded.portal_url,origin_host=excluded.origin_host,
                 portal_kind=excluded.portal_kind,discovery_method=excluded.discovery_method,
                 discovery_evidence=excluded.discovery_evidence,confidence=excluded.confidence,
                 last_seen_at=excluded.last_seen_at,latest_page_status=excluded.latest_page_status,
                 latest_http_status=excluded.latest_http_status,latest_final_url=excluded.latest_final_url,
                 latest_title=excluded.latest_title,
                 latest_jobposting_jsonld=excluded.latest_jobposting_jsonld,active=1""",
            (
                result.website_group_key, portal["portal_url"], portal["canonical_url"],
                canonical_host(portal["canonical_url"]), portal["portal_kind"],
                portal["discovery_method"], portal["discovery_evidence"],
                portal["confidence"], now, now, portal.get("page_status", "not_checked"),
                portal.get("http_status"), portal.get("final_url"),
                portal.get("page_title"), int(portal.get("jobposting_jsonld") or 0),
            ),
        )
    if result.discovery_complete:
        if seen_urls:
            placeholders = ",".join("?" for _ in seen_urls)
            conn.execute(
                """UPDATE institution_recruitment_portals SET active=0
                   WHERE website_group_key=? AND canonical_url NOT IN ("""
                + placeholders + ")",
                [result.website_group_key, *sorted(seen_urls)],
            )
        else:
            conn.execute(
                "UPDATE institution_recruitment_portals SET active=0 WHERE website_group_key=?",
                (result.website_group_key,),
            )
    evidence = {
        "homepage_title": result.homepage_title,
        "candidate_urls": [
            {
                "url": portal.get("canonical_url"),
                "kind": portal.get("portal_kind"),
                "anchor_text": portal.get("anchor_text"),
                "score": portal.get("score"),
                "page_status": portal.get("page_status"),
                "http_status": portal.get("http_status"),
                "final_url": portal.get("final_url"),
                "page_title": portal.get("page_title"),
                "jobposting_jsonld": portal.get("jobposting_jsonld", 0),
                "page_signals": portal.get("page_signals", []),
                "reason": portal.get("reason"),
            }
            for portal in result.portals
        ],
        "rule_version": DISCOVERY_RULE_VERSION,
        "note": "Portal discovery evidence is not vacancy evidence.",
    }
    conn.execute(
        """INSERT INTO institution_recruitment_observations
           (run_id,website_group_key,checked_at,official_website_url,homepage_status,
            homepage_http_status,homepage_final_url,overall_status,candidates_seen,
            portals_found,pages_fetched,request_attempts,discovery_complete,
            candidate_truncated,reason,evidence_json,rule_version)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            run_id, result.website_group_key, result.checked_at, result.official_website_url,
            result.homepage_status, result.homepage_http_status, result.homepage_final_url,
            result.overall_status, result.candidates_seen, len(result.portals),
            result.pages_fetched, result.request_attempts, int(result.discovery_complete),
            int(result.candidate_truncated), result.reason, _json(evidence),
            DISCOVERY_RULE_VERSION,
        ),
    )
    conn.execute(
        """UPDATE institution_recruitment_websites
           SET latest_status=?,latest_checked_at=?,latest_http_status=?,
               latest_portals_found=?,latest_reason=? WHERE website_group_key=?""",
        (
            result.overall_status, result.checked_at, result.homepage_http_status,
            len(result.portals), result.reason, result.website_group_key,
        ),
    )


def scan_registry(registry_db=None, *, max_websites=25, offset=0,
                  include_checked=False, timeout=20, delay=0.75, max_portals=3,
                  only_hosts=None, progress=None):
    """Scan a resumable batch of unique current-scope official websites."""
    registry_path = Path(registry_db or DEFAULT_REGISTRY_DB)
    conn = registry_connect(registry_path)
    ensure_registry_schema(conn)
    bounded = max(1, min(1000, int(max_websites)))
    start_offset = max(0, int(offset))
    conditions = ["w.official_website_url IS NOT NULL", "w.linked_institution_count>0"]
    args = []
    requested_hosts = [str(item).strip() for item in (only_hosts or []) if str(item).strip()]
    if requested_hosts:
        placeholders = ",".join("?" for _ in requested_hosts)
        conditions.append(f"w.website_group_key IN ({placeholders})")
        args.extend(requested_hosts)
    if not include_checked:
        conditions.append(
            "NOT EXISTS (SELECT 1 FROM institution_recruitment_observations o "
            "WHERE o.website_group_key=w.website_group_key AND o.rule_version=?)"
        )
        args.append(DISCOVERY_RULE_VERSION)
    rows = conn.execute(
        """SELECT w.* FROM institution_recruitment_websites AS w WHERE """
        + " AND ".join(conditions)
        + " ORDER BY w.website_group_key LIMIT ? OFFSET?".replace("OFFSET?", "OFFSET ?"),
        args + [bounded, start_offset],
    ).fetchall()
    scope = _current_scope(conn)
    started = now_utc()
    run_id = conn.execute(
        """INSERT INTO registry_runs
           (mode,data_year,seed_scope,started_at,status,requested_count,rule_version)
           VALUES (?,?,?,?,?,?,?)""",
        ("scan", CARNEGIE_DATA_YEAR, scope, started, "running", len(rows),
         DISCOVERY_RULE_VERSION),
    ).lastrowid
    conn.commit()
    reset_robots_cache()
    reset_crawl_budgets()
    scanner = RecruitmentSiteScanner(timeout=timeout, delay=delay)
    counts = {}
    processed = requests = 0
    try:
        for index, row in enumerate(rows, start=1):
            try:
                result = scanner.scan_site(row, max_portals=max_portals)
            except Exception as error:
                status, reason = _exception_status(error)
                result = SiteScanResult(
                    row["website_group_key"], now_utc(), row["official_website_url"], status,
                    None, None, None, status, 0, tuple(), 0, 1, False, False, reason,
                )
            with conn:
                _record_site_scan(conn, result, run_id=run_id)
            processed += 1
            requests += result.request_attempts
            counts[result.overall_status] = counts.get(result.overall_status, 0) + 1
            if progress:
                progress({
                    "index": index,
                    "total": len(rows),
                    "website_group_key": row["website_group_key"],
                    "representative_unitid": row["representative_unitid"],
                    "status": result.overall_status,
                    "portals_found": len(result.portals),
                    "pages_fetched": result.pages_fetched,
                    "reason": result.reason,
                })
        result_summary = {
            "requested": len(rows),
            "processed": processed,
            "request_attempts": requests,
            "statuses": counts,
            "offset": start_offset,
            "include_checked": bool(include_checked),
            "max_portals": max_portals,
            "rule_version": DISCOVERY_RULE_VERSION,
            "seed_scope": scope,
        }
        conn.execute(
            """UPDATE registry_runs SET finished_at=?,status='ok',processed_count=?,
               request_attempts=?,result_counts_json=? WHERE id=?""",
            (now_utc(), processed, requests, _json(result_summary), run_id),
        )
        conn.commit()
        return result_summary
    except Exception as error:
        conn.execute(
            """UPDATE registry_runs SET finished_at=?,status='failed',processed_count=?,
               request_attempts=?,result_counts_json=?,error=? WHERE id=?""",
            (now_utc(), processed, requests, _json(counts), _short(error, 1000), run_id),
        )
        conn.commit()
        raise
    finally:
        conn.close()


def registry_summary(registry_db=None):
    path = Path(registry_db or DEFAULT_REGISTRY_DB)
    if not path.exists():
        return {"exists": False, "registry_db": str(path)}
    conn = registry_connect(path)
    try:
        ensure_registry_schema(conn)
        by_segment = [dict(row) for row in conn.execute(
            """SELECT target_segment,COUNT(*) AS institutions,
                      SUM(official_website_url IS NOT NULL) AS with_official_website
               FROM institution_recruitment_targets GROUP BY target_segment
               ORDER BY target_segment"""
        )]
        by_status = [dict(row) for row in conn.execute(
            """SELECT overall_status AS status,COUNT(*) AS observations
               FROM institution_recruitment_observations GROUP BY overall_status
               ORDER BY status"""
        )]
        return {
            "exists": True,
            "registry_db": str(path),
            "targets": conn.execute(
                "SELECT COUNT(*) FROM institution_recruitment_targets"
            ).fetchone()[0],
            "website_groups": conn.execute(
                "SELECT COUNT(*) FROM institution_recruitment_websites "
                "WHERE linked_institution_count>0"
            ).fetchone()[0],
            "active_portals": conn.execute(
                "SELECT COUNT(*) FROM institution_recruitment_portals WHERE active=1"
            ).fetchone()[0],
            "observations": conn.execute(
                "SELECT COUNT(*) FROM institution_recruitment_observations"
            ).fetchone()[0],
            "by_segment": by_segment,
            "by_status": by_status,
            "scope": _current_scope(conn),
            "rule_version": DISCOVERY_RULE_VERSION,
        }
    finally:
        conn.close()


def export_registry_csv(registry_db=None, output_path=None, *,
                        rule_version=DISCOVERY_RULE_VERSION, overwrite=False):
    """Export one current-scope row per Carnegie target institution."""
    registry_path = Path(registry_db or DEFAULT_REGISTRY_DB)
    target_path = Path(output_path or DEFAULT_EXPORT_CSV)
    if target_path.resolve() == registry_path.resolve():
        raise ValueError("CSV output must not overwrite the registry database")
    if target_path.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists: {target_path}. Pass overwrite=True to replace it."
        )
    conn = registry_connect(registry_path)
    try:
        targets = conn.execute(
            """SELECT t.*,w.linked_institution_count
               FROM institution_recruitment_targets AS t
               LEFT JOIN institution_recruitment_websites AS w
                 ON w.website_group_key=t.website_group_key
               ORDER BY t.target_segment,t.institution_name,t.unitid"""
        ).fetchall()
        latest_observations = {}
        for row in conn.execute(
            """SELECT o.* FROM institution_recruitment_observations AS o
               WHERE o.rule_version=? AND NOT EXISTS (
                 SELECT 1 FROM institution_recruitment_observations AS newer
                 WHERE newer.rule_version=o.rule_version
                   AND newer.website_group_key=o.website_group_key
                   AND (newer.checked_at>o.checked_at OR
                        (newer.checked_at=o.checked_at AND newer.id>o.id))
               )""",
            (rule_version,),
        ):
            latest_observations[row["website_group_key"]] = row
        observation_keys = {
            (row["website_group_key"], row["checked_at"])
            for row in latest_observations.values()
        }
        portals_by_observation = {}
        if observation_keys:
            for row in conn.execute(
                """SELECT website_group_key,last_seen_at,canonical_url,portal_kind,
                          latest_page_status,confidence
                   FROM institution_recruitment_portals
                   ORDER BY website_group_key,canonical_url"""
            ):
                key = (row["website_group_key"], row["last_seen_at"])
                if key in observation_keys:
                    portals_by_observation.setdefault(key, []).append(row)

        fieldnames = [
            "unitid", "institution_name", "city", "state", "control", "data_year",
            "carnegie_research_code", "carnegie_research_designation",
            "carnegie_institutional_code", "carnegie_institutional_classification",
            "is_r2", "is_teaching_oriented", "teaching_proxy_rule",
            "is_business_focused", "target_segment", "source_release_id", "seed_scope",
            "official_website_raw", "official_website_url", "official_host",
            "website_group_key", "linked_institution_count", "latest_status",
            "latest_checked_at", "homepage_status", "homepage_http_status",
            "homepage_final_url", "candidates_seen", "portals_found", "pages_fetched",
            "request_attempts", "discovery_complete", "candidate_truncated", "latest_reason",
            "scan_rule_version", "portal_urls", "portal_kinds", "portal_page_statuses",
            "portal_confidences",
        ]
        target_path.parent.mkdir(parents=True, exist_ok=True)
        rows_written = 0
        portal_keys_written = set()
        portal_attachments_written = 0
        with target_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for target in targets:
                observation = latest_observations.get(target["website_group_key"])
                if target["official_website_url"] is None:
                    status = "missing_official_website"
                elif observation is None:
                    status = "not_scanned"
                else:
                    status = observation["overall_status"]
                portals = []
                if observation is not None:
                    portals = portals_by_observation.get(
                        (observation["website_group_key"], observation["checked_at"]), []
                    )
                portal_attachments_written += len(portals)
                portal_keys_written.update(
                    (observation["website_group_key"], item["canonical_url"])
                    for item in portals if observation is not None
                )

                def joined(column):
                    return " | ".join(str(item[column]) for item in portals if item[column])

                writer.writerow({
                    **{field: target[field] for field in fieldnames
                       if field in target.keys()},
                    "latest_status": status,
                    "latest_checked_at": observation["checked_at"] if observation else None,
                    "homepage_status": observation["homepage_status"] if observation else None,
                    "homepage_http_status": observation["homepage_http_status"] if observation else None,
                    "homepage_final_url": observation["homepage_final_url"] if observation else None,
                    "candidates_seen": observation["candidates_seen"] if observation else None,
                    "portals_found": observation["portals_found"] if observation else None,
                    "pages_fetched": observation["pages_fetched"] if observation else None,
                    "request_attempts": observation["request_attempts"] if observation else None,
                    "discovery_complete": observation["discovery_complete"] if observation else None,
                    "candidate_truncated": observation["candidate_truncated"] if observation else None,
                    "latest_reason": observation["reason"] if observation else None,
                    "scan_rule_version": observation["rule_version"] if observation else None,
                    "portal_urls": joined("canonical_url"),
                    "portal_kinds": joined("portal_kind"),
                    "portal_page_statuses": joined("latest_page_status"),
                    "portal_confidences": joined("confidence"),
                })
                rows_written += 1
        return {
            "registry_db": str(registry_path),
            "output_csv": str(target_path),
            "rows": rows_written,
            "portal_rows": len(portal_keys_written),
            "institution_portal_attachments": portal_attachments_written,
            "observed_groups": len(latest_observations),
            "scope": _current_scope(conn),
            "rule_version": rule_version,
        }
    finally:
        conn.close()


__all__ = [
    "BUSINESS_FOCUS_NAME", "CARNEGIE_DATA_YEAR", "CARNEGIE_PUBLIC_URL",
    "CARNEGIE_RAD_URL", "DEFAULT_EXPORT_CSV", "DEFAULT_PUBLIC_FILE", "DEFAULT_RAD_FILE",
    "DEFAULT_REGISTRY_DB", "DEFAULT_SEED_SCOPE", "DEFAULT_SOURCE_DB",
    "DISCOVERY_RULE_VERSION", "IPEDS_WEBSITE_YEAR", "R2_NAME",
    "build_target_rows", "canonical_host", "download_carnegie_sources",
    "ensure_registry_schema", "export_registry_csv", "is_r2",
    "is_teaching_oriented_classification", "registry_connect", "registry_summary",
    "scan_registry", "seed_registry",
]
