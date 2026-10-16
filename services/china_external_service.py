# -*- coding: utf-8 -*-
"""中国外贸与跨境资金官方数据管线。

来源边界：
- 海关总署英文统计站：进出口当月/累计金额与同比，人民币和美元原始口径；
- 商务部数据中心与官方搜索：实际使用外资、全行业/非金融类对外投资；
- 国家外汇管理局时间序列表：银行结售汇、银行代客涉外收付款及直接投资项。

纪律：不做币种换算；不把银行结售汇或全部跨境收支写成 FDI；缺失保持缺失；
所有写入均幂等，失败不清理旧数据。
"""
from __future__ import annotations

import io
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sqlite3
import threading
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import date, datetime
from decimal import Decimal
from urllib.parse import quote, urljoin

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from services.china_macro_service import connect, ensure_china_macro_tables
from services.safe_account_service import ACCOUNT_INDICATORS, update_safe_accounts
from services.mofcom_fdi_bulletin_service import FDI_ANNUAL_INDICATORS, sync_mofcom_fdi_annual


UA = {'User-Agent': ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                     'AppleWebKit/537.36 Chrome/124 Safari/537.36')}
TIMEOUT = 40

GACC_LIST_URL = 'http://english.customs.gov.cn/Statistics/Statistics'
MOFCOM_SEARCH_URL = 'http://search.mofcom.gov.cn/search/'
MOFCOM_FDI_API = 'https://data.mofcom.gov.cn/datamofcom/front/lywz/direct/query'
MOFCOM_ODI_API = 'https://data.mofcom.gov.cn/datamofcom/front/tzhz/fordirinvest/dateQuery'
MOFCOM_DATA_URL = 'https://data.mofcom.gov.cn/'
MOFCOM_FEC_BASE = 'https://fec.mofcom.gov.cn/'
MOFCOM_FEC_LIST_API = urljoin(MOFCOM_FEC_BASE, 'edi_ecms_web_front/front/fec/list')
MOFCOM_FEC_LIST_PAGE = urljoin(
    MOFCOM_FEC_BASE, 'list/ggfw/tjsj/ydjmtjsj/1/cateinfo.html')
SAFE_SETTLEMENT_PAGE = 'https://www.safe.gov.cn/safe/2023/0215/22329.html'
SAFE_PAYMENT_PAGE = 'https://www.safe.gov.cn/safe/2018/0419/8806.html'
SAFE_RESERVE_INDEX = 'https://www.safe.gov.cn/safe/whcb/'
SAFE_RAW_ROOT = Path(__file__).resolve().parents[1] / 'data' / 'china_macro_raw'


def _new_session():
    """官方站点偶有连接重置；只做有限重试，不绕过证书或验证页。"""
    session = requests.Session()
    retry = Retry(
        total=3, connect=3, read=2, status=2, backoff_factor=0.45,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset(('GET', 'POST')),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry, pool_connections=4, pool_maxsize=4)
    session.mount('http://', adapter)
    session.mount('https://', adapter)
    return session


_THREAD_LOCAL = threading.local()


def _worker_session():
    session = getattr(_THREAD_LOCAL, 'session', None)
    if session is None:
        session = _new_session()
        _THREAD_LOCAL.session = session
    return session


def _meta(name, unit, frequency='monthly'):
    return name, unit, frequency


EXTERNAL_INDICATORS = {}
EXTERNAL_INDICATORS.update(ACCOUNT_INDICATORS)
EXTERNAL_INDICATORS.update(FDI_ANNUAL_INDICATORS)
for currency, unit in (('CNY', '亿元'), ('USD', '亿美元')):
    for subject, label in (('TRADE', '货物进出口总额'), ('EXPORT', '货物出口'),
                           ('IMPORT', '货物进口')):
        EXTERNAL_INDICATORS.update({
            f'CN_GACC_{subject}_{currency}_MONTH': _meta(f'{label}当月({currency})', unit),
            f'CN_GACC_{subject}_{currency}_YTD': _meta(f'{label}累计({currency})', unit),
            f'CN_GACC_{subject}_{currency}_MONTH_YOY': _meta(f'{label}当月同比({currency})', '%'),
            f'CN_GACC_{subject}_{currency}_YTD_YOY': _meta(f'{label}累计同比({currency})', '%'),
        })
    EXTERNAL_INDICATORS.update({
        f'CN_GACC_BALANCE_{currency}_MONTH': _meta(f'货物贸易差额当月({currency})', unit),
        f'CN_GACC_BALANCE_{currency}_YTD': _meta(f'货物贸易差额累计({currency})', unit),
        f'CN_MOFCOM_FDI_{currency}_YTD': _meta(f'实际使用外资累计({currency})', unit),
        f'CN_MOFCOM_FDI_{currency}_YTD_YOY': _meta(f'实际使用外资累计同比({currency})', '%'),
        f'CN_MOFCOM_ODI_ALL_{currency}_YTD': _meta(f'全行业对外直接投资累计({currency})', unit),
        f'CN_MOFCOM_ODI_ALL_{currency}_YTD_YOY': _meta(f'全行业对外直接投资累计同比({currency})', '%'),
        f'CN_MOFCOM_ODI_NONFIN_{currency}_YTD': _meta(f'非金融类对外直接投资累计({currency})', unit),
        f'CN_MOFCOM_ODI_NONFIN_{currency}_YTD_YOY': _meta(f'非金融类对外直接投资累计同比({currency})', '%'),
        f'CN_SAFE_FX_SETTLEMENT_{currency}': _meta(f'银行结汇当月({currency})', unit),
        f'CN_SAFE_FX_SALES_{currency}': _meta(f'银行售汇当月({currency})', unit),
        f'CN_SAFE_FX_NET_{currency}': _meta(f'银行结售汇差额当月({currency})', unit),
        f'CN_SAFE_XB_RECEIPTS_{currency}': _meta(f'银行代客涉外收入当月({currency})', unit),
        f'CN_SAFE_XB_PAYMENTS_{currency}': _meta(f'银行代客涉外支出当月({currency})', unit),
        f'CN_SAFE_XB_NET_{currency}': _meta(f'银行代客涉外收支差额当月({currency})', unit),
        f'CN_SAFE_DI_RECEIPTS_{currency}': _meta(f'银行代客直接投资项收入当月({currency})', unit),
        f'CN_SAFE_DI_PAYMENTS_{currency}': _meta(f'银行代客直接投资项支出当月({currency})', unit),
        f'CN_SAFE_DI_NET_{currency}': _meta(f'银行代客直接投资项差额当月({currency})', unit),
    })

EXTERNAL_INDICATORS.update({
    'CN_MOFCOM_FDI_NEW_FIRMS_YTD': _meta('新设外商投资企业累计', '家'),
    'CN_MOFCOM_FDI_NEW_FIRMS_YTD_YOY': _meta('新设外商投资企业累计同比', '%'),
    'CN_MOFCOM_ODI_ENTERPRISES_YTD': _meta('境外非金融类投资企业累计', '家'),
    # 国家外汇管理局《官方储备资产表》：与 PBoC 资产负债表中的
    # foreign_exchange、银行结售汇和涉外收付款分开保存，不能互相替代。
    'CN_SAFE_FX_RESERVES_USD': _meta('官方储备资产：外汇储备', '亿美元'),
    'CN_SAFE_IMF_RESERVE_POSITION_USD': _meta('官方储备资产：在基金组织的储备头寸', '亿美元'),
    'CN_SAFE_SDR_HOLDINGS_USD': _meta('官方储备资产：特别提款权', '亿美元'),
    'CN_SAFE_MONETARY_GOLD_USD': _meta('官方储备资产：黄金（美元计价）', '亿美元'),
    'CN_SAFE_MONETARY_GOLD_QUANTITY': _meta('官方储备资产：黄金数量', '万盎司'),
    'CN_SAFE_OTHER_RESERVE_ASSETS_USD': _meta('官方储备资产：其他储备资产', '亿美元'),
    'CN_SAFE_OFFICIAL_RESERVE_ASSETS_TOTAL_USD': _meta('官方储备资产合计', '亿美元'),
})

SAFE_RESERVE_INDICATOR_CODES = (
    'CN_SAFE_FX_RESERVES_USD',
    'CN_SAFE_IMF_RESERVE_POSITION_USD',
    'CN_SAFE_SDR_HOLDINGS_USD',
    'CN_SAFE_MONETARY_GOLD_USD',
    'CN_SAFE_MONETARY_GOLD_QUANTITY',
    'CN_SAFE_OTHER_RESERVE_ASSETS_USD',
    'CN_SAFE_OFFICIAL_RESERVE_ASSETS_TOTAL_USD',
)


def _number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(float(value)) else None
    text = str(value).strip().replace(',', '').replace('，', '')
    if text in ('', '-', '—', '–', '…', '...'):
        return None
    match = re.search(r'-?\d+(?:\.\d+)?', text)
    return float(match.group(0)) if match else None


def _period_yyyymm(value):
    text = re.sub(r'\D', '', str(value or ''))
    if len(text) >= 6 and 1980 <= int(text[:4]) <= 2100 and 1 <= int(text[4:6]) <= 12:
        return f'{text[:4]}-{text[4:6]}'
    return None


def _signed(direction, value):
    number = _number(value)
    if number is None:
        return None
    return -abs(number) if direction in ('下降', '减少', '回落') else abs(number)


def _observation(code, period, value, source_name, source_type, source_url,
                 source_title, parser_notes, data_status='official', formula=None,
                 release_date=None, retrieved_at=None, raw_sha256=None,
                 parser_version=None, missing_reason=None,
                 source_page_created_date=None,
                 source_page_last_updated_date=None,
                 release_date_not_provided=False, release_date_reason=None):
    if code not in EXTERNAL_INDICATORS or value is None or not period:
        return None
    name, unit, frequency = EXTERNAL_INDICATORS[code]
    return {
        'indicator_code': code, 'indicator_name': name, 'period': period,
        'value': float(value), 'unit': unit, 'frequency': frequency,
        'data_status': data_status, 'source_name': source_name,
        'source_type': source_type, 'source_url': source_url,
        'source_title': source_title, 'parser_notes': parser_notes,
        'formula': formula,
        'release_date': release_date,
        'retrieved_at': retrieved_at or datetime.now().isoformat(),
        'raw_sha256': raw_sha256,
        'parser_version': parser_version or 'china_external_v1',
        'missing_reason': missing_reason,
        'source_page_created_date': source_page_created_date,
        'source_page_last_updated_date': source_page_last_updated_date,
        'release_date_not_provided': int(bool(release_date_not_provided)),
        'release_date_reason': release_date_reason,
    }


def upsert_external_observations(conn, rows, commit=True):
    now, count = datetime.now().isoformat(), 0
    for row in rows:
        if not row:
            continue
        cur = conn.execute('''INSERT INTO china_macro_observations (
            indicator_code,indicator_name,period,value,unit,frequency,data_status,
            source_name,source_type,source_url,source_title,parser_notes,formula,
            release_date,retrieved_at,raw_sha256,parser_version,missing_reason,
            source_page_created_date,source_page_last_updated_date,
            release_date_not_provided,release_date_reason,updated_at
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(indicator_code,period) DO UPDATE SET
            indicator_name=excluded.indicator_name,value=excluded.value,
            unit=excluded.unit,frequency=excluded.frequency,
            data_status=excluded.data_status,source_name=excluded.source_name,
            source_type=excluded.source_type,source_url=excluded.source_url,
            source_title=excluded.source_title,parser_notes=excluded.parser_notes,
            formula=excluded.formula,release_date=excluded.release_date,
            retrieved_at=excluded.retrieved_at,raw_sha256=excluded.raw_sha256,
            parser_version=excluded.parser_version,missing_reason=excluded.missing_reason,
            source_page_created_date=excluded.source_page_created_date,
            source_page_last_updated_date=excluded.source_page_last_updated_date,
            release_date_not_provided=excluded.release_date_not_provided,
            release_date_reason=excluded.release_date_reason,
            updated_at=excluded.updated_at''', (
            row['indicator_code'], row['indicator_name'], row['period'], row['value'],
            row['unit'], row['frequency'], row['data_status'], row['source_name'],
            row['source_type'], row['source_url'], row['source_title'],
            row['parser_notes'], row.get('formula'), row.get('release_date'),
            row.get('retrieved_at') or now, row.get('raw_sha256'),
            row.get('parser_version') or 'china_external_v1',
            row.get('missing_reason'), row.get('source_page_created_date'),
            row.get('source_page_last_updated_date'),
            int(bool(row.get('release_date_not_provided'))),
            row.get('release_date_reason'), now))
        count += cur.rowcount
    if commit:conn.commit()
    return count


# ── 海关总署：双币种贸易 ──────────────────────────────────────────────────────

def parse_gacc_q3_government_relay(html, period, source_url):
    """Read the inspected national September/YTD CNY statements, never USD or balances.

    This is an explicitly labelled government relay, not the GACC original table.
    The article's stated2026 data year and publication metadata govern the period.
    """
    if period!='2026-09' or source_url!='https://gxt.hunan.gov.cn/gxt/xxgk_71033/gzdt/rdjj/202610/t20261015_34081504.html':
        raise ValueError('Use the reviewed national2026Q3 government-relay profile.')
    soup=BeautifulSoup(html,'html.parser');content=soup.select_one('#j-show-body')
    published=soup.find('meta',attrs={'name':'PubDate'})
    if content is None or published is None:raise ValueError('The exact article body and publication date are required.')
    published_text=published.get('content','')
    if not re.match(r'^2026-10-\d{2}',published_text):raise ValueError('The relay publication year/month does not corroborate the data period.')
    text=re.sub(r'\s+','',unicodedata.normalize('NFKC',content.get_text(' ',strip=True)))
    if ('海关总署' not in text or '国新办' not in text
            or '介绍2026年前三季度我国货物贸易进出口情况' not in text
            or '10月14日' not in text):
        raise ValueError('Explicit national scope, data year and release-event date are required.')
    cumulative=re.search(r'据海关统计,今年前三季度,我国货物贸易([^。]+)。',text)
    monthly=re.search(r'9月份进出口([^。]+)。其中,([^。]+)。',text)
    if not cumulative or not monthly:raise ValueError('Do not bridge cumulative and monthly statements or infer missing components.')
    amount=r'([\d,]+(?:\.\d+)?)(万亿元|亿元)'
    change=r'(?:同比)?(增长|下降)(\d+(?:\.\d+)?)%'
    trade_ytd=re.search(r'今年前三季度我国货物贸易进出口'+amount+','+change,text)
    if not trade_ytd:raise ValueError('One explicit national cumulative total and its published rate are required.')
    specs=[('TRADE','YTD',trade_ytd)]
    for subject,label in [('EXPORT','出口'),('IMPORT','进口')]:
        match=re.search(label+amount+','+change,cumulative[1])
        if not match:raise ValueError('A cumulative component is absent: '+subject)
        specs.append((subject,'YTD',match))
    total=re.search(r'^'+amount+r',[^。]*?,'+change,monthly[1])
    if not total:raise ValueError('One explicit monthly total and rate are required.')
    specs.append(('TRADE','MONTH',total))
    for subject,label in [('EXPORT','出口'),('IMPORT','进口')]:
        match=re.search(label+amount+','+change,monthly[2])
        if not match:raise ValueError('A monthly component is absent: '+subject)
        specs.append((subject,'MONTH',match))
    rows=[]
    for subject,kind,match in specs:
        value=float(Decimal(match[1].replace(',',''))*(Decimal(10000) if match[2]=='万亿元' else Decimal(1)))
        rate=float(match[4])*(-1 if match[3]=='下降' else 1)
        notes=('政府官网转引海关发布会数据，非海关原始总值表；当月与累计分别取原文；'
               '金额原单位='+match[2]+'，仅按金额量纲转为亿元，未做汇率换算；保留新闻稿公布精度；'
               '同比为原文公布值，不由两年金额直除；relay_published_at='+published_text+'; statement='+match[0])
        for code,number in [(f'CN_GACC_{subject}_CNY_{kind}',value),(f'CN_GACC_{subject}_CNY_{kind}_YOY',rate)]:
            row=_observation(code,period,number,'海关总署（湖南省工信厅官网转引）','gacc_trade',source_url,
                '9月进出口首破5万亿元 装备制造业贡献近9成出口增量',notes,
                data_status='official_relayed',release_date='2026-10-14',parser_version='gacc-government-relay-20261015',
                source_page_created_date='2026-10-15' if soup.find('meta',attrs={'name':'others','content':re.compile('页面生成时间 2026-10-15')}) else None,
                release_date_reason='资料发布会日期据转载正文10月14日及明确2026资料年；转载发布日与生成日另记，不冒充已取得海关原稿。')
            rows.append(row)
    values={r['indicator_code']:r['value'] for r in rows}
    for kind in ('MONTH','YTD'):
        if not math.isclose(values[f'CN_GACC_TRADE_CNY_{kind}'],values[f'CN_GACC_EXPORT_CNY_{kind}']+values[f'CN_GACC_IMPORT_CNY_{kind}'],abs_tol=100,rel_tol=0):
            raise ValueError('Published rounded total and components do not reconcile at news precision.')
    return rows

_MONTHS = {name: index for index, name in enumerate(
           ('jan', 'feb', 'mar', 'apr', 'may', 'jun',
            'jul', 'aug', 'sep', 'oct', 'nov', 'dec'), 1)}


def parse_gacc_title(title):
    text = re.sub(r'\s+', ' ', title or '').strip()
    combined = re.search(
        r"Total Export\s*&\s*Import Values,\s*(?:January|Jan\.?)[-–—]"
        r"(?:February|Feb\.?)\s*(20\d{2}).*?\(in\s*(CNY|USD)\)", text, re.I)
    if combined:
        return f'{combined.group(1)}-02', combined.group(2).upper()
    match = re.search(r"Total Export\s*&\s*Import Values,\s*([A-Za-z]+)\.?\s*(20\d{2}).*?\(in\s*(CNY|USD)\)",
                      text, re.I)
    if not match:
        return None, None
    month = _MONTHS.get(match.group(1)[:3].lower())
    return (f'{match.group(2)}-{month:02d}', match.group(3).upper()) if month else (None, None)


def parse_gacc_article(html, title='', source_url=''):
    soup = BeautifulSoup(html or '', 'html.parser')
    page_title = title or (soup.find('tr').get_text(' ', strip=True) if soup.find('tr') else '')
    period, currency = parse_gacc_title(page_title)
    if not period:
        return []
    labels = {
        'Total Export & Import': 'TRADE', 'Total Export': 'EXPORT',
        'Total Import': 'IMPORT', 'Export-Import Balance': 'BALANCE',
    }
    rows = []
    for tr in soup.find_all('tr'):
        cells = [re.sub(r'\s+', ' ', cell.get_text(' ', strip=True)).strip()
                 for cell in tr.find_all(['td', 'th'])]
        if not cells or cells[0] not in labels:
            continue
        subject = labels[cells[0]]
        values = [_number(value) for value in cells[1:]]
        if subject == 'BALANCE':
            pairs = ((f'CN_GACC_BALANCE_{currency}_MONTH', values[0] if values else None),
                     (f'CN_GACC_BALANCE_{currency}_YTD', values[1] if len(values) > 1 else None))
        else:
            pairs = (
                (f'CN_GACC_{subject}_{currency}_MONTH', values[0] if values else None),
                (f'CN_GACC_{subject}_{currency}_YTD', values[1] if len(values) > 1 else None),
                (f'CN_GACC_{subject}_{currency}_MONTH_YOY', values[3] if len(values) > 3 else None),
                (f'CN_GACC_{subject}_{currency}_YTD_YOY', values[4] if len(values) > 4 else None),
            )
        for code, value in pairs:
            row = _observation(
                code, period, value, '海关总署', 'gacc_trade', source_url,
                page_title, '海关总署月度总值表原始币种；未做汇率换算')
            if row:
                rows.append(row)
    return rows


def _fetch_gacc_listing_page(page, session=None):
    active = session or _worker_session()
    response = active.get(GACC_LIST_URL, params={'page': page}, headers=UA, timeout=TIMEOUT)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or 'utf-8'
    return page, response.text


def _parse_gacc_listing_page(html):
    links, periods = [], []
    for anchor in BeautifulSoup(html, 'html.parser').find_all('a', href=True):
        title = re.sub(r'\s+', ' ', anchor.get_text(' ', strip=True)).strip()
        if not re.match(r"^\(1\)\s*China's Total Export\s*&\s*Import Values", title, re.I):
            continue
        period, currency = parse_gacc_title(title)
        if period:
            periods.append(period)
            links.append((period, currency, urljoin(GACC_LIST_URL, anchor['href']), title))
    return periods, links


def discover_gacc_releases(start_year, max_pages=180, sleep_seconds=0.12, session=None):
    session = session or _new_session()
    found, empty_pages = {}, 0
    parallel = max_pages > 60
    page = 1
    executor = ThreadPoolExecutor(max_workers=3) if parallel else None
    try:
        while page <= max_pages:
            batch = list(range(page, min(max_pages + 1, page + (9 if parallel else 1))))
            if parallel:
                fetched = executor.map(lambda number: _fetch_gacc_listing_page(number), batch)
            else:
                fetched = [_fetch_gacc_listing_page(batch[0], session)]
            should_stop = False
            for _, html in sorted(fetched):
                page_periods, links = _parse_gacc_listing_page(html)
                for period, currency, url, title in links:
                    if int(period[:4]) >= int(start_year):
                        found[(period, currency)] = (url, title)
                empty_pages = empty_pages + 1 if not page_periods else 0
                if page_periods and max(int(period[:4]) for period in page_periods) < int(start_year):
                    should_stop = True
                if empty_pages >= 3:
                    should_stop = True
            if should_stop:
                break
            page = batch[-1] + 1
            time.sleep(sleep_seconds)
    finally:
        if executor:
            executor.shutdown(wait=True)
    return found


def update_gacc_trade(conn, start_year, max_pages=180, sleep_seconds=0.12, session=None):
    session = session or _new_session()
    errors, parsed = [], []
    try:
        releases = discover_gacc_releases(start_year, max_pages, sleep_seconds, session)
    except Exception as exc:
        return {'records_upserted': 0, 'releases': 0, 'errors': [f'discover: {exc}']}
    release_items = [(url, title) for (_, _), (url, title) in sorted(releases.items())]

    def fetch(item):
        url, title = item
        try:
            active = _worker_session() if len(release_items) > 60 else session
            response = active.get(url, headers=UA, timeout=TIMEOUT)
            response.raise_for_status()
            response.encoding = response.apparent_encoding or 'utf-8'
            return parse_gacc_article(response.text, title, url), None
        except Exception as exc:
            return [], f'{title}: {exc}'

    if len(release_items) > 60:
        with ThreadPoolExecutor(max_workers=3) as executor:
            results = executor.map(fetch, release_items)
            for rows, error in results:
                parsed.extend(rows)
                if error:
                    errors.append(error)
    else:
        for item in release_items:
            rows, error = fetch(item)
            parsed.extend(rows)
            if error:
                errors.append(error)
            time.sleep(sleep_seconds)
    return {'records_upserted': upsert_external_observations(conn, parsed),
            'releases': len(releases), 'observations_parsed': len(parsed), 'errors': errors}


# ── 国家外汇管理局：结售汇与涉外收付款时间序列 ───────────────────────────────

def _discover_safe_attachment(page_url, link_keyword, session=None):
    session = session or _new_session()
    response = session.get(page_url, headers=UA, timeout=TIMEOUT)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or 'utf-8'
    for anchor in BeautifulSoup(response.text, 'html.parser').find_all('a', href=True):
        text = re.sub(r'\s+', '', anchor.get_text(' ', strip=True))
        if link_keyword in text and '时间序列' in text:
            return urljoin(page_url, anchor['href'])
    raise ValueError(f'未发现{link_keyword}时间序列附件')


def _excel_month(value, datemode=0):
    if isinstance(value, datetime):
        return value.strftime('%Y-%m')
    if isinstance(value, date):
        return value.strftime('%Y-%m')
    if isinstance(value, (int, float)):
        try:
            import xlrd
            return xlrd.xldate_as_datetime(float(value), datemode).strftime('%Y-%m')
        except Exception:
            return None
    text = str(value or '').strip()
    match = re.search(r'(20\d{2})[年./-]?(1[0-2]|0?[1-9])', text)
    return f'{match.group(1)}-{int(match.group(2)):02d}' if match else None


def parse_safe_wide_rows(headers, value_rows, code_by_row, source_name, source_type,
                         source_url, source_title, datemode=0):
    """解析 SAFE 横向月份表；公开为纯函数，测试可用小型列表而非二进制工作簿。"""
    observations = []
    for column, raw_date in enumerate(headers):
        period = _excel_month(raw_date, datemode)
        if not period:
            continue
        for row_key, values in value_rows.items():
            value = _number(values[column]) if column < len(values) else None
            code = code_by_row.get(row_key)
            row = _observation(
                code, period, value, source_name, source_type, source_url, source_title,
                '外汇局官方月度时间序列表原值；差额采用官方表中差额行，未做币种换算')
            if row:
                observations.append(row)
    return observations


def _parse_safe_settlement_workbook(content, attachment_url):
    import openpyxl
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    observations = []
    for currency, sheet_name in (('CNY', '以人民币计价（月度）'),
                                 ('USD', '以美元计价（月度）')):
        sheet = workbook[sheet_name]
        raw = [list(row) for row in sheet.iter_rows(min_row=1, max_row=45,
                                                     max_col=420, values_only=True)]
        labels = {str(row[0]).strip(): index for index, row in enumerate(raw) if row and row[0]}
        header_index = next(index for index, row in enumerate(raw)
                            if row and str(row[0]).strip() == '项目')
        headers = raw[header_index]
        row_names = {'settlement': '一、结汇', 'sales': '二、售汇', 'net': '三、差额'}
        value_rows = {key: raw[labels[label]] for key, label in row_names.items()}
        codes = {
            'settlement': f'CN_SAFE_FX_SETTLEMENT_{currency}',
            'sales': f'CN_SAFE_FX_SALES_{currency}',
            'net': f'CN_SAFE_FX_NET_{currency}',
        }
        observations.extend(parse_safe_wide_rows(
            headers, value_rows, codes, '国家外汇管理局', 'safe_fx_settlement',
            attachment_url, '银行结售汇数据时间序列'))
    return observations


def _parse_safe_payment_workbook(content, attachment_url):
    import xlrd
    workbook = xlrd.open_workbook(file_contents=content)
    observations = []
    for currency, sheet_name in (('CNY', '以人民币计值（月度）'),
                                 ('USD', '以美元计值（月度）')):
        sheet = workbook.sheet_by_name(sheet_name)
        matrix = [sheet.row_values(index) for index in range(sheet.nrows)]
        header_index = next(index for index, row in enumerate(matrix)
                            if row and str(row[0]).strip() == '项目')
        top_rows, direct_rows = {}, []
        for index, row in enumerate(matrix):
            label = re.sub(r'\s+', '', str(row[0] if row else ''))
            if label == '一、收入':
                top_rows['receipts'] = row
            elif label == '二、支出':
                top_rows['payments'] = row
            elif label == '三、收支差额':
                top_rows['net'] = row
            elif label in ('其中：直接投资', '其中:直接投资'):
                direct_rows.append(row)
        if len(direct_rows) < 3 or len(top_rows) < 3:
            raise ValueError(f'{sheet_name} 行标签结构不完整')
        value_rows = {
            **top_rows, 'di_receipts': direct_rows[0],
            'di_payments': direct_rows[1], 'di_net': direct_rows[2],
        }
        codes = {
            'receipts': f'CN_SAFE_XB_RECEIPTS_{currency}',
            'payments': f'CN_SAFE_XB_PAYMENTS_{currency}',
            'net': f'CN_SAFE_XB_NET_{currency}',
            'di_receipts': f'CN_SAFE_DI_RECEIPTS_{currency}',
            'di_payments': f'CN_SAFE_DI_PAYMENTS_{currency}',
            'di_net': f'CN_SAFE_DI_NET_{currency}',
        }
        observations.extend(parse_safe_wide_rows(
            matrix[header_index], value_rows, codes, '国家外汇管理局',
            'safe_cross_border_payments', attachment_url,
            '银行代客涉外收付款数据时间序列', workbook.datemode))
    return observations


def update_safe_flows(conn, session=None):
    session = session or _new_session()
    errors, rows, attachments = [], [], {}
    jobs = (
        ('settlement', SAFE_SETTLEMENT_PAGE, '银行结售汇数据', _parse_safe_settlement_workbook),
        ('payments', SAFE_PAYMENT_PAGE, '银行代客涉外收付款数据', _parse_safe_payment_workbook),
    )
    for key, page_url, keyword, parser in jobs:
        try:
            attachment = _discover_safe_attachment(page_url, keyword, session)
            response = session.get(attachment, headers=UA, timeout=TIMEOUT)
            response.raise_for_status()
            attachments[key] = attachment
            rows.extend(parser(response.content, attachment))
        except Exception as exc:
            errors.append(f'{key}: {exc}')
    return {'records_upserted': upsert_external_observations(conn, rows),
            'observations_parsed': len(rows), 'attachments': attachments, 'errors': errors}


# ── 国家外汇管理局：官方储备资产 ────────────────────────────────────────────

def _safe_reserve_cell(cell):
    """Read SAFE's HTML table cell without treating a dash as zero.

    The current official table exposes the numeric value in the ``x:num``
    attribute.  Text is retained as a fallback for older/translated pages.
    """
    if cell is None:
        return None
    raw = cell.get('x:num') or cell.get('data-value')
    if raw is None:
        raw = cell.get_text(' ', strip=True)
    return _number(raw)


def _save_safe_raw(content, source_url):
    SAFE_RAW_ROOT.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(content).hexdigest()
    suffix = '.xlsx' if source_url.lower().endswith('.xlsx') else ('.xls' if source_url.lower().endswith('.xls') else '.html')
    stamp = datetime.now().strftime('%Y%m%dT%H%M%S%f')
    path = SAFE_RAW_ROOT / f'safe_reserve_{stamp}{suffix}'
    part = path.with_suffix(path.suffix + '.part')
    part.write_bytes(content)
    os.replace(part, path)
    return str(path.relative_to(SAFE_RAW_ROOT.parent.parent)).replace(os.sep, '/'), digest


def _safe_reserve_label(row):
    return re.sub(r'\s+', ' ', row[0].get_text(' ', strip=True)).strip() if row else ''


def _safe_page_path_date(source_url):
    """Return a URL path date as page metadata, never as a release date.

    SAFE keeps one annual reserve-assets page and updates its monthly columns
    in place.  The path date therefore describes the page location/creation,
    not the release date of every observation later added to that table.
    """
    match = re.search(r'/safe/(20\d{2})/(\d{2})(\d{2})/', str(source_url or ''))
    return (f'{match.group(1)}-{match.group(2)}-{match.group(3)}'
            if match else None)


def _safe_release_date(_source_url, _source_title=''):
    """Return no per-observation date unless SAFE exposes direct evidence.

    This compatibility helper intentionally does not parse URL path dates.
    The annual page/listing dates are stored in dedicated source-page fields.
    """
    return None


def parse_safe_reserve_html(html, source_url, source_title='',
                            source_page_last_updated_date=None):
    """Parse one SAFE annual official-reserve-assets HTML table.

    Each month occupies two adjacent columns (USD and SDR).  This parser only
    records the first USD column and the separate gold-quantity row; it never
    converts SDR to USD and never invents future-month observations.
    """
    soup = BeautifulSoup(html or '', 'html.parser')
    table = next((table for table in soup.find_all('table')
                  if 'Foreign currency reserves' in table.get_text(' ', strip=True)
                  or '官方储备资产' in table.get_text(' ', strip=True)), None)
    if table is None:
        raise ValueError('SAFE official reserve-assets table not found')
    table_rows = [row.find_all(['td', 'th']) for row in table.find_all('tr')]
    month_names = {
        'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'may': 5, 'jun': 6,
        'jul': 7, 'aug': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dec': 12,
    }
    month_row = next((row for row in table_rows
                      if sum(1 for cell in row
                             if re.match(r'^20\d{2}[./-]\d{1,2}$',
                                         cell.get_text(' ', strip=True))
                             or re.sub(r'\s+', ' ', cell.get_text(' ', strip=True)).lower()[:3] in month_names) >= 2), None)
    if month_row is None:
        raise ValueError('SAFE reserve-assets month header not found')
    month_columns = []
    for index, cell in enumerate(month_row):
        label = re.sub(r'\s+', ' ', cell.get_text(' ', strip=True))
        numeric_month = re.match(r'^(20\d{2})[./-](\d{1,2})$', label)
        if numeric_month:
            month_columns.append((index, int(numeric_month.group(1)), int(numeric_month.group(2))))
            continue
        key = label.lower()[:3]
        if key in month_names:
            month_columns.append((index, None, month_names[key]))
    # Current SAFE HTML uses one month label per cell but two value cells per
    # month (USD followed by SDR). Older fixtures/pages may put the month
    # labels directly above each pair; preserve that direct mapping there.
    header_has_sdr_labels = any('sdr' in cell.get_text(' ', strip=True).lower()
                                for cell in month_row)
    if len(month_columns) >= 6 and not header_has_sdr_labels:
        month_columns = [(1 + 2 * ordinal, header_year, month)
                         for ordinal, (_header_col, header_year, month)
                         in enumerate(month_columns)]
    year_match = re.search(r'(20\d{2})', f'{source_title} {source_url}')
    year = int(year_match.group(1)) if year_match else None
    labels = {}
    quantity_row = None
    gold_row_index = None
    for row_index, row in enumerate(table_rows):
        label = _safe_reserve_label(row)
        compact = re.sub(r'\s+', ' ', label).lower()
        if 'foreign currency reserves' in compact and 'official' not in compact:
            labels['fx'] = row
        elif 'imf reserve position' in compact:
            labels['imf'] = row
        # Use plural SDRs so the footnote's singular "SDR" does not overwrite
        # the actual row with an empty note row.
        elif re.search(r'\bsdrs\b|特别提款权', compact) and '折算汇率' not in compact:
            labels['sdr'] = row
        elif ('quantity' in compact and 'gold' in compact) or '黄金数量' in compact:
            quantity_row = row
        elif re.search(r'\bgold\b|黄金', compact):
            labels['gold'] = row
            gold_row_index = row_index
        elif 'other reserve assets' in compact or '其他储备资产' in compact:
            labels['other'] = row
        elif 'total' in compact or '合计' in compact:
            labels['total'] = row
    # Historical HTML pages sometimes leave the quantity label blank. Accept
    # only an immediately adjacent row with an explicit ounce unit; a blank
    # numeric row without that unit is not evidence for inventing a quantity.
    if quantity_row is None and gold_row_index is not None:
        for candidate in table_rows[gold_row_index + 1:gold_row_index + 3]:
            text = ' '.join(cell.get_text(' ', strip=True) for cell in candidate)
            if re.search(r'万\s*盎司|fine\s+troy\s+ounce|ounce', text, re.I):
                quantity_row = candidate
                break
    code_map = {
        'fx': 'CN_SAFE_FX_RESERVES_USD',
        'imf': 'CN_SAFE_IMF_RESERVE_POSITION_USD',
        'sdr': 'CN_SAFE_SDR_HOLDINGS_USD',
        'gold': 'CN_SAFE_MONETARY_GOLD_USD',
        'other': 'CN_SAFE_OTHER_RESERVE_ASSETS_USD',
        'total': 'CN_SAFE_OFFICIAL_RESERVE_ASSETS_TOTAL_USD',
    }
    if 'total' not in labels:
        raise ValueError('SAFE reserve-assets total row not found')
    rows = []
    raw_digest = hashlib.sha256((html or '').encode('utf-8')).hexdigest()
    retrieved_at = datetime.now().isoformat()
    release_date = _safe_release_date(source_url, source_title)
    page_created_date = _safe_page_path_date(source_url)
    release_reason = (
        'SAFE annual reserve-assets page does not expose a distinct '
        'per-observation release date; URL path date is page metadata only.')
    for column, header_year, month in month_columns:
        period_year = header_year or year
        if period_year is None or not 1 <= month <= 12:
            continue
        period = f'{period_year:04d}-{month:02d}'
        for key, code in code_map.items():
            row = labels.get(key)
            # The first column is the label; the first monthly USD value is
            # therefore at the next cell after the header label.
            value_column = column if column < len(row or []) else column + 1
            value = _safe_reserve_cell(row[value_column] if row and value_column < len(row) else None)
            item = _observation(code, period, value, '国家外汇管理局',
                                'safe_official_reserve_assets', source_url,
                                source_title or f'官方储备资产 {year}',
                                'SAFE 官方储备资产年度 HTML 表；单位为亿美元；未将 SDR 换算为美元',
                                release_date=release_date, retrieved_at=retrieved_at,
                                raw_sha256=raw_digest, parser_version='safe_reserve_v3',
                                source_page_created_date=page_created_date,
                                source_page_last_updated_date=source_page_last_updated_date,
                                release_date_not_provided=True,
                                release_date_reason=release_reason)
            if item:
                rows.append(item)
        if quantity_row:
            value_column = column if column < len(quantity_row) else column + 1
            value = _safe_reserve_cell(quantity_row[value_column]
                                       if value_column < len(quantity_row) else None)
            item = _observation('CN_SAFE_MONETARY_GOLD_QUANTITY', period, value,
                                '国家外汇管理局', 'safe_official_reserve_assets', source_url,
                                source_title or f'官方储备资产 {year}',
                                'SAFE 官方储备资产表中的黄金数量；单位为万盎司；不与美元金额相加',
                                release_date=release_date, retrieved_at=retrieved_at,
                                raw_sha256=raw_digest, parser_version='safe_reserve_v3',
                                source_page_created_date=page_created_date,
                                source_page_last_updated_date=source_page_last_updated_date,
                                release_date_not_provided=True,
                                release_date_reason=release_reason)
            if item:
                rows.append(item)
    return rows


def parse_safe_reserve_workbook(content, source_url, source_title='',
                                source_page_last_updated_date=None):
    """Parse SAFE's pre-2018 XLS/XLSX reserve-assets attachments.

    Older official releases use one month per column (2015) or paired USD/SDR
    columns (2016-2017). The parser keeps the first USD column and does not
    convert the SDR companion column.
    """
    matrixes = []
    if content[:2] == b'PK':
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        matrixes = [[list(row) for row in sheet.iter_rows(values_only=True)]
                    for sheet in workbook.worksheets]
    else:
        import xlrd
        workbook = xlrd.open_workbook(file_contents=content)
        matrixes = [[sheet.row_values(index) for index in range(sheet.nrows)]
                    for sheet in (workbook.sheet_by_index(i) for i in range(workbook.nsheets))]
    rows = []
    raw_digest = hashlib.sha256(content).hexdigest()
    retrieved_at = datetime.now().isoformat()
    release_date = _safe_release_date(source_url, source_title)
    page_created_date = _safe_page_path_date(source_url)
    release_reason = (
        'SAFE historical reserve-assets attachment does not expose a distinct '
        'per-observation release date in the parsed table; URL date was not inferred.')
    for matrix in matrixes:
        header_index = next((i for i, row in enumerate(matrix)
                             if sum(1 for cell in row if re.match(
                                 r'^20\d{2}[./-]\d{1,2}$', str(cell).strip())
                                 or (isinstance(cell, (int, float)) and 2000 <= float(cell) <= 2100)) >= 2), None)
        if header_index is None:
            continue
        month_columns = []
        for index, cell in enumerate(matrix[header_index]):
            # Excel may expose 2017.10 as the float 2017.1. Formatting numeric
            # cells to two decimals preserves October instead of misreading it
            # as January.
            text = f'{float(cell):.2f}' if isinstance(cell, (int, float)) else str(cell).strip()
            match = re.match(r'^(20\d{2})[./-](\d{1,2})$', text)
            if match:
                month_columns.append((index, int(match.group(1)), int(match.group(2))))
        if not month_columns:
            continue
        label_rows = {}
        quantity_rows = []
        for row_index, row in enumerate(matrix[header_index + 1:], header_index + 1):
            label_parts = [str(row[0] if row else '')]
            if row_index + 1 < len(matrix):
                # One historical XLSX places the English label on the next
                # row while keeping the numeric values on the preceding row.
                label_parts.append(str(matrix[row_index + 1][0] or ''))
            label = re.sub(r'\s+', ' ', ' '.join(label_parts)).lower()
            if not label:
                # Some XLSX sheets place the English label on the next row;
                # quantity rows are identified from their cell values below.
                label = ' '.join(str(cell) for cell in row if cell is not None).lower()
            has_month_value = any(_number(row[column] if column < len(row) else None) is not None
                                  for column, _year, _month in month_columns)
            if 'foreign currency reserves' in label and has_month_value:
                label_rows['fx'] = row
            elif 'imf reserve position' in label and has_month_value:
                label_rows['imf'] = row
            elif re.search(r'\bsdrs\b', label) and has_month_value:
                label_rows['sdr'] = row
            elif re.search(r'\bgold\b', label) and has_month_value:
                label_rows['gold'] = row
            elif 'other reserve assets' in label and has_month_value:
                label_rows['other'] = row
            elif 'total' in label and has_month_value:
                label_rows['total'] = row
            if any('万盎司' in str(cell) or 'fine troy' in str(cell).lower()
                   for cell in row):
                quantity_rows.append(row)
        if 'total' not in label_rows:
            continue
        code_map = {
            'fx': 'CN_SAFE_FX_RESERVES_USD',
            'imf': 'CN_SAFE_IMF_RESERVE_POSITION_USD',
            'sdr': 'CN_SAFE_SDR_HOLDINGS_USD',
            'gold': 'CN_SAFE_MONETARY_GOLD_USD',
            'other': 'CN_SAFE_OTHER_RESERVE_ASSETS_USD',
            'total': 'CN_SAFE_OFFICIAL_RESERVE_ASSETS_TOTAL_USD',
        }
        for column, year, month in month_columns:
            period = f'{year:04d}-{month:02d}'
            for key, code in code_map.items():
                row = label_rows.get(key)
                value = _number(row[column] if row and column < len(row) else None)
                item = _observation(code, period, value, '国家外汇管理局',
                                    'safe_official_reserve_assets', source_url,
                                    source_title or f'官方储备资产 {year}',
                                    'SAFE 官方 XLS/XLSX 历史附件；单位为亿美元；只读取第一美元列',
                                    release_date=release_date, retrieved_at=retrieved_at,
                                    raw_sha256=raw_digest, parser_version='safe_reserve_v3',
                                    source_page_created_date=page_created_date,
                                    source_page_last_updated_date=source_page_last_updated_date,
                                    release_date_not_provided=True,
                                    release_date_reason=release_reason)
                if item:
                    rows.append(item)
            if quantity_rows:
                quantity = _number(quantity_rows[0][column]
                                   if column < len(quantity_rows[0]) else None)
                item = _observation('CN_SAFE_MONETARY_GOLD_QUANTITY', period, quantity,
                                    '国家外汇管理局', 'safe_official_reserve_assets', source_url,
                                    source_title or f'官方储备资产 {year}',
                                    'SAFE 官方 XLS/XLSX 历史附件中的黄金数量；单位为万盎司',
                                    release_date=release_date, retrieved_at=retrieved_at,
                                    raw_sha256=raw_digest, parser_version='safe_reserve_v3',
                                    source_page_created_date=page_created_date,
                                    source_page_last_updated_date=source_page_last_updated_date,
                                    release_date_not_provided=True,
                                    release_date_reason=release_reason)
                if item:
                    rows.append(item)
        if rows:
            break
    return rows


def _safe_reserve_articles(session, full_history=False):
    response = session.get(SAFE_RESERVE_INDEX, headers=UA, timeout=TIMEOUT)
    response.raise_for_status()
    # SAFE pages are UTF-8 even when the apparent-encoding detector reports a
    # legacy Chinese codec; using the wrong codec hides the official article
    # title and the HTML attachment link.
    response.encoding = 'utf-8'
    found = {}
    for anchor in BeautifulSoup(response.text, 'html.parser').find_all('a', href=True):
        url = urljoin(SAFE_RESERVE_INDEX, anchor['href'])
        year_match = re.search(r'/safe/(20\d{2})/', url)
        if not year_match or not url.endswith('.html'):
            continue
        title = re.sub(r'\s+', ' ', anchor.get_text(' ', strip=True))
        if '官方储备资产' not in title and 'official reserve assets' not in title.lower():
            continue
        listing_context = (anchor.parent.parent.get_text(' ', strip=True)
                           if anchor.parent and anchor.parent.parent else
                           anchor.parent.get_text(' ', strip=True)
                           if anchor.parent else title)
        listing_dates = re.findall(r'20\d{2}-\d{2}-\d{2}', listing_context)
        page_last_updated_date = listing_dates[-1] if listing_dates else None
        # The index item often opens a metadata page; follow its explicit
        # ``html`` attachment link to the table page, still within SAFE.
        detail = session.get(url, headers=UA, timeout=TIMEOUT)
        detail.raise_for_status()
        detail.encoding = 'utf-8'
        detail_soup = BeautifulSoup(detail.text, 'html.parser')
        html_link = next((item for item in detail_soup.find_all('a', href=True)
                          if item.get_text(' ', strip=True).lower() == 'html'), None)
        if html_link:
            actual_url = urljoin(url, html_link['href'])
            title_years = re.findall(r'20\d{2}', title)
            key_year = int(title_years[0]) if title_years else int(year_match.group(1))
            found[key_year] = {
                'url': actual_url, 'title': title,
                'listing_url': url,
                'page_created_date': _safe_page_path_date(actual_url),
                'page_last_updated_date': page_last_updated_date,
            }
        else:
            # The 2015-2017 release exposes three official spreadsheet
            # attachments rather than an HTML table. Keep each attachment as
            # its own source item so its year/month columns remain auditable.
            for attachment in detail_soup.find_all('a', href=True):
                attachment_url = urljoin(url, attachment['href'])
                attachment_title = re.sub(r'\s+', ' ', attachment.get_text(' ', strip=True))
                if not attachment_url.lower().endswith(('.xls', '.xlsx')):
                    continue
                years = re.findall(r'20\d{2}', attachment_title)
                if years:
                    found[int(years[0])] = {
                        'url': attachment_url, 'title': attachment_title,
                        'listing_url': url,
                        'page_created_date': _safe_page_path_date(url),
                        'page_last_updated_date': page_last_updated_date,
                    }
    if not found:
        raise ValueError('SAFE official reserve-assets article list not found')
    if not full_history:
        floor = datetime.now().year - 1
        found = {year: item for year, item in found.items() if year >= floor}
    return sorted(found.items())


def update_safe_reserves(conn, full_history=False, session=None):
    """Fetch SAFE official reserve-assets articles and upsert only parsed rows."""
    session = session or _new_session()
    rows, errors, articles = [], [], []
    latest_raw_digest = None
    latest_period = None
    try:
        articles = _safe_reserve_articles(session, full_history)
    except Exception as exc:
        return {'records_upserted': 0, 'observations_parsed': 0,
                'articles': [], 'errors': [str(exc)]}
    for year, article in articles:
        url = article['url']
        title = article['title']
        page_last_updated_date = article.get('page_last_updated_date')
        try:
            response = session.get(url, headers=UA, timeout=TIMEOUT)
            response.raise_for_status()
            response.encoding = 'utf-8'
            raw_path, raw_digest = _save_safe_raw(response.content, url)
            latest_raw_digest = raw_digest
            if url.lower().endswith(('.xls', '.xlsx')):
                parsed = parse_safe_reserve_workbook(
                    response.content, url, title or str(year),
                    source_page_last_updated_date=page_last_updated_date)
            else:
                parsed = parse_safe_reserve_html(
                    response.text, url, title or str(year),
                    source_page_last_updated_date=page_last_updated_date)
            for row in parsed:
                row['parser_notes'] = (row.get('parser_notes') or '') + \
                    f'; raw_snapshot={raw_path}; raw_sha256={raw_digest}'
                row['raw_sha256'] = raw_digest
                row['retrieved_at'] = datetime.now().isoformat()
                row['parser_version'] = 'safe_reserve_v3'
                if row.get('value') is not None:
                    latest_period = max(latest_period or row['period'], row['period'])
            rows.extend(parsed)
        except Exception as exc:
            errors.append(f'{year}: {exc}')
        time.sleep(0.08)
    records_upserted = upsert_external_observations(conn, rows)
    # Preserve existing blank future rows as explicit non-release records. The
    # parser deliberately skips blank cells; this audit update prevents an old
    # placeholder from remaining indistinguishable from an untracked failure.
    if latest_period:
        conn.execute('''UPDATE china_macro_observations
            SET data_status='not_released',
                missing_reason='official_source_future_month_blank',
                parser_version=COALESCE(parser_version, 'safe_reserve_v2'),
                retrieved_at=COALESCE(retrieved_at, ?),
                raw_sha256=COALESCE(raw_sha256, ?),
                updated_at=?
            WHERE indicator_code IN ({}) AND value IS NULL
              AND period > ?'''.format(','.join('?' for _ in SAFE_RESERVE_INDICATOR_CODES)),
                     (datetime.now().isoformat(), latest_raw_digest,
                      datetime.now().isoformat(), *SAFE_RESERVE_INDICATOR_CODES,
                      latest_period))
        conn.commit()
    return {'records_upserted': records_upserted,
            'observations_parsed': len(rows),
            'articles': [article['url'] for _, article in articles],
            'errors': errors}


# ── 商务部：吸收外资与对外直接投资 ────────────────────────────────────────────

def _mofcom_period(text):
    compact = re.sub(r'\s+', '', text or '')
    match = re.search(
        r'(20\d{2})年(?:(?:1|一)[—\-至](\d{1,2})月|前(\d{1,2})个月|'
        r'(\d{1,2})月|(一季度|上半年|前三季度|全年))', compact)
    if not match:
        # “2025年全国吸收外资……”这类不带月份的年度发布才回落到 12 月；
        # 不能用可选月份正则，否则“一季度”会先被裸年份吞掉并误记为全年。
        match = re.search(r'(20\d{2})年(?=[，,:：]?(?:全年)?(?:全国|我国).{0,18}'
                          r'(?:吸收外资|实际使用外资|外商投资|全行业对外直接投资))',
                          compact)
        return f'{match.group(1)}-12' if match else None
    quarter_months = {'一季度': 3, '上半年': 6, '前三季度': 9, '全年': 12}
    month = int(match.group(2) or match.group(3) or match.group(4) or
                quarter_months.get(match.group(5), 12))
    return f'{match.group(1)}-{month:02d}'


def parse_mofcom_fdi_text(title, summary='', source_url=''):
    title_period, body_period = _mofcom_period(title), _mofcom_period(summary)
    # A headline's statistical year must not be replaced by a later comparison
    # such as "较2022年全年" inside a 2023 annual release.
    period = title_period or body_period
    title_compact = re.sub(r'\s+', '', title or '')
    if (title_period and body_period and title_period[:4] == body_period[:4]
            and title_period.endswith('-12') and not body_period.endswith('-12')
            and not re.search(r'全年|\d{1,2}月|季度|半年|前\d+个月', title_compact)):
        period = body_period
    text = re.sub(r'\s+', '', summary or '')
    title_text = re.sub(r'\s+', '', title or '')
    # 年末发布常在同一篇正文先列全年累计、后列“12月当月”。YTD 指标只能读取
    # 当月段之前的累计句；否则 2017 年会把 8775.6 亿元误写成当月 739.4 亿元。
    current_month = re.search(r'(?:\d{1,2}月)?当月', text)
    cumulative_text = text[:current_month.start()] if current_month else text
    values = {}
    amount = re.search(
        r'(?:实际使用外资(?:金额)?|全国吸收外资)([\d.]+)亿元(?:人民币)?',
        cumulative_text)
    if not amount:
        amount = re.search(
            r'(?:实际使用外资(?:金额)?|全国吸收外资)([\d.]+)亿元(?:人民币)?',
            title_text)
    if amount:
        values['CN_MOFCOM_FDI_CNY_YTD'] = float(amount.group(1))
    yoy = re.search(
        r'实际使用外资(?:金额)?[\d.]+亿元(?:人民币)?，?(?:按可比口径)?(?:同比|比上年)'
        r'(增长|下降)([\d.]+)%', cumulative_text)
    if yoy:
        values['CN_MOFCOM_FDI_CNY_YTD_YOY'] = _signed(yoy.group(1), yoy.group(2))
    # Only a directly printed USD clause is accepted; never derive it from CNY.
    usd = re.search(r'(?:折合|以美元计(?:为)?)([\d.]+)亿美元，?(?:同比)?'
                    r'(增长|下降)([\d.]+)%', cumulative_text)
    if usd:
        values['CN_MOFCOM_FDI_USD_YTD'] = float(usd[1])
        values['CN_MOFCOM_FDI_USD_YTD_YOY'] = _signed(usd[2], usd[3])
    firms = re.search(
        r'新设立外商投资企业(\d+)家，?(?:同比|比上年)(增长|下降)([\d.]+)%',
        cumulative_text)
    if firms:
        values['CN_MOFCOM_FDI_NEW_FIRMS_YTD'] = float(firms.group(1))
        values['CN_MOFCOM_FDI_NEW_FIRMS_YTD_YOY'] = _signed(firms.group(2), firms.group(3))
    return [_observation(code, period, value, '商务部', 'mofcom_fdi', source_url,
                         title, '商务部吸收外资官方发布；累计口径')
            for code, value in values.items() if value is not None and period]


def parse_mofcom_odi_text(title, summary='', source_url=''):
    period = _mofcom_period(summary) or _mofcom_period(title)
    text = re.sub(r'\s+', '', summary or '')
    values = {}
    all_match = re.search(
        r'全行业对外直接投资([\d.]+)亿元人民币，?(?:同比|比上年)(增长|下降)([\d.]+)%'
        r'（以美元计为([\d.]+)亿美元，?(增长|下降)([\d.]+)%）', text)
    if all_match:
        values.update({
            'CN_MOFCOM_ODI_ALL_CNY_YTD': float(all_match.group(1)),
            'CN_MOFCOM_ODI_ALL_CNY_YTD_YOY': _signed(all_match.group(2), all_match.group(3)),
            'CN_MOFCOM_ODI_ALL_USD_YTD': float(all_match.group(4)),
            'CN_MOFCOM_ODI_ALL_USD_YTD_YOY': _signed(all_match.group(5), all_match.group(6)),
        })
    nonfin = re.search(
        r'非金融类直接投资，?累计投资([\d.]+)亿元人民币，?(增长|下降)([\d.]+)%'
        r'（以美元计为([\d.]+)亿美元，?(增长|下降)([\d.]+)%）', text)
    if nonfin:
        values.update({
            'CN_MOFCOM_ODI_NONFIN_CNY_YTD': float(nonfin.group(1)),
            'CN_MOFCOM_ODI_NONFIN_CNY_YTD_YOY': _signed(nonfin.group(2), nonfin.group(3)),
            'CN_MOFCOM_ODI_NONFIN_USD_YTD': float(nonfin.group(4)),
            'CN_MOFCOM_ODI_NONFIN_USD_YTD_YOY': _signed(nonfin.group(5), nonfin.group(6)),
        })
    return [_observation(code, period, value, '商务部、国家外汇管理局', 'mofcom_odi',
                         source_url, title, '商务部对外投资官方发布；累计口径')
            for code, value in values.items() if value is not None]


def _mofcom_search_records(query_text, max_pages, session=None, sleep_seconds=0.12):
    session = session or _new_session()
    records, seen = [], set()
    for page in range(1, max_pages + 1):
        response = session.get(MOFCOM_SEARCH_URL, params={
            'keyWordType': 'title', 'acSuggest': query_text, 'page': page,
        }, headers=UA, timeout=TIMEOUT)
        response.raise_for_status()
        response.encoding = response.apparent_encoding or 'utf-8'
        soup = BeautifulSoup(response.text, 'html.parser')
        page_count = 0
        for item in soup.select('li'):
            anchor = item.select_one('a.a2[href]') or item.select_one('.title a[href]')
            if not anchor:
                continue
            title = re.sub(r'\s+', ' ', anchor.get_text(' ', strip=True)).strip()
            if not title or not _mofcom_period(title):
                continue
            url = urljoin(MOFCOM_SEARCH_URL, anchor['href'])
            summary_node = item.select_one('.summary')
            summary = re.sub(r'\s+', ' ', summary_node.get_text(' ', strip=True)).strip() if summary_node else ''
            key = (url, title, summary)
            if key not in seen:
                seen.add(key)
                records.append({'title': title, 'summary': summary, 'url': url})
                page_count += 1
        if page_count == 0 and page > 1:
            break
        time.sleep(sleep_seconds)
    return records


def _mofcom_fec_odi_records(years, session=None, sleep_seconds=0.08):
    """从商务部对外投资合作司官方列表补最新 ODI 发布。

    该列表接口的中文标题目前会被服务端替换成问号，因此只使用其 URL/发布日期，
    再读取文章正文并用完整统计句筛选；不从问号标题推断任何数值。
    """
    session = session or _new_session()
    records, seen = [], set()
    headers = {**UA, 'Referer': MOFCOM_FEC_LIST_PAGE}
    for year in years:
        page, max_page = 1, 1
        while page <= max_page:
            response = session.post(MOFCOM_FEC_LIST_API, data={
                'pageNumber': page, 'cid': 'ggfw/tjsj/ydjmtjsj', 'year': int(year),
            }, headers=headers, timeout=TIMEOUT)
            response.raise_for_status()
            payload = response.json()
            info = payload.get('pageInfo') or {}
            max_page = int(info.get('maxPageNum') or 1)
            for raw in info.get('rows') or []:
                url = urljoin(MOFCOM_FEC_BASE, raw.get('url') or '')
                if not url or url in seen:
                    continue
                seen.add(url)
                try:
                    article = session.get(url, headers=UA, timeout=TIMEOUT)
                    article.raise_for_status()
                    article.encoding = article.apparent_encoding or 'utf-8'
                    text = re.sub(r'\s+', ' ', BeautifulSoup(
                        article.text, 'html.parser').get_text(' ', strip=True))
                    if '全行业对外直接投资' not in text or '非金融类直接投资' not in text:
                        continue
                    compact = re.sub(r'\s+', '', text)
                    match = re.search(r'(20\d{2}年(?:1[—\-至]\d{1,2}月|\d{1,2}月|上半年|全年)?)，?我国全行业对外直接投资',
                                      compact)
                    title = ((match.group(1) if match else str(year) + '年') +
                             '我国全行业对外直接投资简明统计')
                    records.append({'title': title, 'summary': text, 'url': url})
                except Exception:
                    # 单篇失败由上层搜索/API互补；列表继续，不清旧记录。
                    continue
                time.sleep(sleep_seconds)
            page += 1
    return records


def _mofcom_api_rows(session=None):
    session = session or _new_session()
    headers = {**UA, 'Referer': MOFCOM_DATA_URL}
    rows, errors = [], []
    try:
        payload = session.post(MOFCOM_FDI_API, json={}, headers=headers, timeout=TIMEOUT).json()
        for raw in payload[0] if payload and isinstance(payload[0], list) else []:
            period = _period_yyyymm(raw.get('data_time'))
            pairs = (
                ('CN_MOFCOM_FDI_USD_YTD', _number(raw.get('actual_amount'))),
                ('CN_MOFCOM_FDI_USD_YTD_YOY', _number(raw.get('actual_amount_per'))),
                ('CN_MOFCOM_FDI_NEW_FIRMS_YTD', _number(raw.get('project_number'))),
                ('CN_MOFCOM_FDI_NEW_FIRMS_YTD_YOY', _number(raw.get('project_per'))),
            )
            for code, value in pairs:
                item = _observation(code, period, value, '商务部', 'mofcom_data_api',
                                    MOFCOM_DATA_URL, '吸收外商直接投资月报图表',
                                    '商务部数据中心官方接口；美元累计值按原表保存')
                if item:
                    rows.append(item)
    except Exception as exc:
        errors.append(f'fdi api: {exc}')
    try:
        payload = session.post(MOFCOM_ODI_API, json={}, headers=headers, timeout=TIMEOUT).json()
        for raw in payload[0] if payload and isinstance(payload[0], list) else []:
            period = _period_yyyymm(raw.get('time'))
            pairs = (
                ('CN_MOFCOM_ODI_NONFIN_USD_YTD', _number(raw.get('non_financial'))),
                ('CN_MOFCOM_ODI_NONFIN_USD_YTD_YOY', _number(raw.get('yoy'))),
                ('CN_MOFCOM_ODI_ENTERPRISES_YTD', _number(raw.get('num_enterprice'))),
            )
            for code, value in pairs:
                item = _observation(code, period, value, '商务部', 'mofcom_data_api',
                                    MOFCOM_DATA_URL, '非金融类对外直接投资月报图表',
                                    '商务部数据中心官方接口；美元累计值按原表保存')
                if item:
                    rows.append(item)
    except Exception as exc:
        errors.append(f'odi api: {exc}')
    return rows, errors


def quarantine_mofcom_period_mismatches(conn, rows):
    """Retain complete old rows before excluding an independently verified wrong period."""
    conn.execute('''CREATE TABLE IF NOT EXISTS mofcom_period_corrections (
        id INTEGER PRIMARY KEY AUTOINCREMENT, previous_row_json TEXT NOT NULL,
        correct_period TEXT NOT NULL, correction_reason TEXT NOT NULL, corrected_at TEXT NOT NULL)''')
    expected = {(row['indicator_code'], row['source_url']): row['period'] for row in rows
                if row and row.get('source_type') in ('mofcom_fdi', 'mofcom_odi')
                and row.get('source_url') not in (MOFCOM_DATA_URL, '')}
    corrected = 0
    for (code, url), period in expected.items():
        aliases = {url, url.replace('http://', 'https://', 1), url.replace('https://', 'http://', 1)}
        placeholders = ','.join('?' for _ in aliases)
        old_rows = conn.execute('SELECT * FROM china_macro_observations WHERE indicator_code=? '
                                f'AND source_url IN ({placeholders}) AND period<>? AND value IS NOT NULL',
                                (code, *sorted(aliases), period)).fetchall()
        for old in old_rows:
            reason = f'Official release reporting period is {period}; original row and value retained.'
            conn.execute('INSERT INTO mofcom_period_corrections '
                         '(previous_row_json,correct_period,correction_reason,corrected_at) VALUES (?,?,?,?)',
                         (json.dumps(dict(old), ensure_ascii=False), period, reason,
                          datetime.now().astimezone().isoformat()))
            conn.execute('UPDATE china_macro_observations SET original_value=COALESCE(original_value,value),'
                         "value=NULL,data_status='invalidated_source_period',missing_reason=? WHERE id=?",
                         (reason, old['id']))
            corrected += 1
    return corrected


def update_mofcom_investment(conn, full_history=False, session=None):
    session = session or _new_session()
    rows, errors = _mofcom_api_rows(session)
    jobs = (
        ('全国吸收外资', 18 if full_history else 4, parse_mofcom_fdi_text),
        ('全行业对外直接投资简明统计', 12 if full_history else 4, parse_mofcom_odi_text),
    )
    for query_text, pages, parser in jobs:
        try:
            for record in _mofcom_search_records(query_text, pages, session):
                summary = record['summary']
                if parser is parse_mofcom_fdi_text:
                    try:
                        article = session.get(record['url'], headers=UA, timeout=TIMEOUT)
                        article.raise_for_status()
                        article.encoding = article.apparent_encoding or 'utf-8'
                        summary = BeautifulSoup(article.text, 'html.parser').get_text(' ', strip=True)
                    except Exception:
                        # 搜索摘要仍可作为降级输入；解析器会在“当月”之前找累计句。
                        summary = summary or ''
                rows.extend(parser(record['title'], summary, record['url']))
        except Exception as exc:
            errors.append(f'{query_text}: {exc}')
    try:
        current_year = datetime.now().year
        fec_years = range(2021 if full_history else current_year - 1, current_year + 1)
        for record in _mofcom_fec_odi_records(fec_years, session):
            rows.extend(parse_mofcom_odi_text(
                record['title'], record['summary'], record['url']))
    except Exception as exc:
        errors.append(f'对外投资合作司列表: {exc}')
    # 一篇商务部官方文章只对应一个累计统计期。若此前把发布日期年份或季度误作
    # 全年，仅清理同 code + 同 source_url 的错误期重复；绝不按模块清空旧数据。
    corrected = quarantine_mofcom_period_mismatches(conn, rows)
    conn.commit()
    count = upsert_external_observations(conn, rows)
    annual = sync_mofcom_fdi_annual(conn)
    return {'records_upserted': count + annual['records_changed'],
            'annual_bulletin': annual,
            'observations_parsed': len(rows), 'period_rows_corrected': corrected,
            'errors': errors}


def update_china_external(db_path, full_history=False, start_year=None, sleep_seconds=0.12):
    """更新三组外部官方源。默认回查近两年海关页；full_history 回填到 start_year。"""
    started = datetime.now().isoformat()
    start = int(start_year or (2015 if full_history else datetime.now().year - 1))
    max_gacc_pages = 360 if full_history else 30
    session = _new_session()
    with closing(connect(db_path)) as conn:
        ensure_china_macro_tables(conn)
        safe = update_safe_flows(conn, session)
        safe_reserves = update_safe_reserves(conn, full_history=full_history, session=session)
        mofcom = update_mofcom_investment(conn, full_history, session)
        gacc = update_gacc_trade(conn, start, max_gacc_pages, sleep_seconds, session)
    accounts=update_safe_accounts(db_path,session=session)
    total = sum(item.get('records_upserted', 0) for item in
                (safe, safe_reserves, mofcom, gacc, accounts))
    errors = (safe.get('errors', []) + safe_reserves.get('errors', []) +
              mofcom.get('errors', []) + gacc.get('errors', []) +
              [f'SAFE accounts: {item}' for item in accounts.get('errors',[])])
    return {
        'success': total > 0 or not errors, 'started_at': started,
        'finished_at': datetime.now().isoformat(), 'records_upserted': total,
        'start_year': start, 'full_history': bool(full_history),
        'safe': safe, 'safe_reserves': safe_reserves,
        'safe_accounts': accounts,
        'mofcom': mofcom, 'gacc': gacc, 'errors': errors,
    }
