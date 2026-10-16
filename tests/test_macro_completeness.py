import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook
from services.china_macro_service import parse_cpi_article, parse_fai_article, parse_pmi_article, parse_econ_article, econ_period_from_article, _econ_period, parse_econ_unemployment_observations
from services.china_macro_extended_service import (
    profit_period, capacity_period, parse_industrial_profit_article, parse_annual_livelihood_article, parse_annual_unemployment_commentary,
)
from services.macro_data_quality_service import calendar_coverage, build_macro_data_quality
from services.pboc_balance_sheet_service import parse_balance_sheet_xlsx, latest_balance_vintages
from services.us_macro_service import derive_change_series, derive_difference_series, SERIES, build_yield_curve_snapshots, NOMINAL_TENORS
from services.bank_supervision_snapshot_service import (
    sync_bank_supervision_snapshot, reviewed_history_rows, SNAPSHOT_PATH, HISTORY_PATH,
)
from services.china_macro_service import connect, ensure_china_macro_tables, update_china_macro
from scripts.backfill_macro_gaps import nbs_article_release_date
from scripts.audit_bank_workbook import reported_cell_value, reviewed_workbook_cells, ROW_LABELS
from scripts.capture_macro_history_handoff import safe_link


class MacroCompletenessTests(unittest.TestCase):
    def test_inline_source_links_do_not_expose_local_paths_or_query_parameters(self):
        self.assertEqual(safe_link('https://www.stats.gov.cn/source.html'),
                         'https://www.stats.gov.cn/source.html')
        for url in ('file:///D:/private/data.csv', 'https://example.org/data?version=1',
                    'https://name:secret@example.org/data', 'https://example.org/data#fragment', None):
            self.assertIsNone(safe_link(url))

    def test_bank_workbook_display_units_and_half_up_rounding(self):
        self.assertEqual(reported_cell_value(0.01688, '0.00%'), 1.69)
        self.assertEqual(reported_cell_value(0.00815, '0.00%'), 0.82)
        self.assertEqual(reported_cell_value(0.78755, '0.00%'), 78.76)
        self.assertEqual(reported_cell_value(6594.526443835, '0_ '), 6595)
        self.assertEqual(reported_cell_value(0, '0.00%'), 0)
        self.assertIsNone(reported_cell_value(None, '0.00%'))
        for value, number_format in ((True, '0'), (float('nan'), '0'), ('=A1', '0'), (1, 'General')):
            with self.assertRaises(ValueError):
                reported_cell_value(value, number_format)

    def test_bank_workbook_reconciliation_validates_year_units_and_quarters(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'bank-source.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = '商业银行季度'
            sheet['A1'] = '商业银行主要监管指标情况表（2022年）'
            sheet['E2'] = '单位：亿元、%'
            for column, quarter in enumerate(('一季度', '二季度', '三季度', '四季度'), 2):
                sheet.cell(3, column, quarter)
            for row_number, label in enumerate(ROW_LABELS.values(), 6):
                sheet.cell(row_number, 1, label)
                for column in range(2, 6):
                    sheet.cell(row_number, column, 1).number_format = '0'
            workbook.save(path)
            self.assertEqual(len(reviewed_workbook_cells(path, 2022)), 52)
            with self.assertRaisesRegex(ValueError, 'statistical year'):
                reviewed_workbook_cells(path, 2023)
            sheet['B3'] = '全年'
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'quarter columns'):
                reviewed_workbook_cells(path, 2022)
            sheet['B3'] = '一季度'
            sheet['E2'] = '单位：万元'
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'units'):
                reviewed_workbook_cells(path, 2022)
            sheet['E2'] = '单位：亿元、%'
            sheet['B6'] = None
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'blank'):
                reviewed_workbook_cells(path, 2022)
            workbook.close()

    def test_original_nbs_dated_heading_is_not_a_migrated_url_or_document_date(self):
        self.assertEqual(nbs_article_release_date(
            '<p>一季度国民经济实现良好开局（2018年4月17日）中华人民共和国国家统计局</p>'),
            '2018-04-17')
        self.assertIsNone(nbs_article_release_date(
            '<p>成文日期2018年4月17日；2018年3月份及一季度主要统计数据</p>'))
        self.assertIsNone(nbs_article_release_date(
            '<meta name="PubDate" content="2023-02-03"><p>（2018年4月17日）中华人民共和国国家统计局</p>'))

    def test_source_publication_date_is_not_the_reporting_year(self):
        html='<meta name="PubDate" content="2026/01/19 10:00"><p>2025年12月份及全年主要统计数据</p>'
        self.assertEqual(nbs_article_release_date(html),'2026-01-19')
        self.assertEqual(econ_period_from_article(html),'2025-12')

    def test_missing_conflicting_or_invalid_publication_metadata_stays_unknown(self):
        self.assertIsNone(nbs_article_release_date('2026年7月15日；6月份调查失业率5.0%'))
        self.assertIsNone(nbs_article_release_date('<meta name="PubDate" content="2026/02/31">'))
        self.assertIsNone(nbs_article_release_date('<meta name="PubDate" content="2026/01/19"><meta name="PubDate" content="2026/01/20">'))

    def test_annual_commentary_only_uses_explicit_paired_months_not_average(self):
        html='2024年，全国城镇调查失业率平均值为5.1%。1、2月份冬季影响，失业率升至5.2%、5.3%。31个大城市调查失业率为5.0%。'
        rows=parse_annual_unemployment_commentary(html,2024)
        self.assertEqual(rows,[('2024-01','CN_UNEMP_SURVEY',5.2),('2024-02','CN_UNEMP_SURVEY',5.3)])
        self.assertEqual(parse_annual_unemployment_commentary('2024年，全国城镇调查失业率平均值为5.1%。',2024),[])
        self.assertEqual(parse_annual_unemployment_commentary(html,2023),[])

    def test_monthly_commentary_explicit_pairs_and_title_year(self):
        html='<h2>解读2019年1—2月就业形势</h2><p>今年1、2月份全国城镇调查失业率分别为5.1%、5.3%，同比分别上升0.1、0.3个百分点。</p>'
        self.assertEqual(parse_annual_unemployment_commentary(html,2019),
            [('2019-01','CN_UNEMP_SURVEY',5.1),('2019-02','CN_UNEMP_SURVEY',5.3)])
        self.assertEqual(parse_annual_unemployment_commentary(html,2020),[])

    def test_annual_month_range_list_and_single_month_have_exact_values(self):
        html='<p>2023年，全国城镇调查失业率平均值为5.2%。分月看，1—2月份失业率延续上年年底态势，处于5.5%、5.6%的水平。3月份，随着生产生活恢复，失业率降至5.3%。</p><p>全年31个大城市失业率平均值为5.1%。1月份失业率为5.8%。</p>'
        self.assertEqual(parse_annual_unemployment_commentary(html,2023),
            [('2023-01','CN_UNEMP_SURVEY',5.5),('2023-02','CN_UNEMP_SURVEY',5.6),('2023-03','CN_UNEMP_SURVEY',5.3)])

    def test_rate_range_and_month_range_average_are_not_allocated(self):
        html='<p>2023年，全国城镇调查失业率平均值为5.2%。1—2月份失业率在5.5%—5.6%的区间。4—6月份失业率平均为5.2%。</p>'
        self.assertEqual(parse_annual_unemployment_commentary(html,2023),[])

    def test_month_specific_rate_wordings_do_not_fall_back_to_another_month(self):
        html='2020年9月份主要统计数据。9月份全国城镇调查失业率是5.4%。3月份全国城镇调查失业率为5.9%。'
        self.assertEqual(parse_econ_article(html,period='2020-09')['CN_UNEMP_SURVEY'],5.4)
        self.assertNotIn('CN_UNEMP_SURVEY',parse_econ_article(html,period='2020-04'))

    def test_national_subject_omitted_in_explicit_month_pair_keeps_subgroups_separate(self):
        html='<p>上半年，全国城镇调查失业率平均为5.7%。4月份，全国城镇调查失业率为6.1%；5、6月份连续回落，分别为5.9%、5.5%。6月份，本地户籍人口调查失业率为5.3%。</p>'
        self.assertEqual(parse_econ_article(html,period='2022-06')['CN_UNEMP_SURVEY'],5.5)
        self.assertNotIn('CN_UNEMP_SURVEY',parse_econ_article(html,period='2022-03'))

    def test_conflicting_same_month_and_other_population_are_not_reported(self):
        html='<p>2023年，全国城镇调查失业率平均值为5.2%。3月份失业率为5.3%。3月份失业率为5.4%。4月份青年失业率为20.4%。</p>'
        self.assertEqual(parse_annual_unemployment_commentary(html,2023),[])

    def test_default_target_rate_cannot_override_explicit_month_conflict(self):
        html = ('<p>1月份全国城镇调查失业率为5.1%。</p>'
                '<p>2月份全国城镇调查失业率为5.2%。2月份全国城镇调查失业率为5.3%。</p>')
        self.assertEqual(parse_econ_unemployment_observations(html, '2025-02'), [
            ('2025-01', 'CN_UNEMP_SURVEY', 5.1)])
    def test_h15_curve_uses_supported_tenors_and_same_date_without_forward_fill(self):
        self.assertEqual(len(NOMINAL_TENORS),11)
        self.assertNotIn('DGS2MO',{row[0] for row in SERIES})
        rows=build_yield_curve_snapshots({'DGS1MO':[{'period':'2026-10-07','value':4.07}],
              'DGS10':[{'period':'2026-10-07','value':5.28},{'period':'2026-10-08','value':5.22}]})
        self.assertEqual(rows[-1]['period'],'2026-10-08')
        self.assertNotIn('DGS1MO',rows[-1]['values'])
        self.assertEqual(rows[0]['values']['DGS1MO'],4.07)

    def test_combined_month_quarter_header_uses_actual_month_not_migration_date(self):
        url='https://www.stats.gov.cn/sj/zxfb/202302/t20230203_1899988.html'
        self.assertEqual(econ_period_from_article('2018年6月份及上半年主要统计数据','上半年国民经济总体平稳',url),'2018-06')
        self.assertEqual(econ_period_from_article('2022年3月份及一季度主要统计数据','一季度国民经济开局总体平稳',url),'2022-03')
        self.assertIsNone(econ_period_from_article('2022年5月份及上半年主要统计数据','上半年国民经济总体平稳',url))

    def test_explicit_january_unemployment_is_not_merged_flow_imputation(self):
        html='1月份，全国城镇调查失业率为5.1%。2月份，全国城镇调查失业率为5.2%。31个大城市城镇调查失业率为5.3%。'
        rows=parse_econ_unemployment_observations(html,'2025-02')
        self.assertIn(('2025-01','CN_UNEMP_SURVEY',5.1),rows)
        self.assertIn(('2025-02','CN_UNEMP_SURVEY',5.2),rows)

    def test_explicit_jan_feb_list_is_not_split_from_an_average(self):
        html='<p>1月和2月份，全国城镇调查失业率分别为5.3%和6.2%，环比分别上升0.1和0.9个百分点。</p>'
        rows=parse_econ_unemployment_observations(html,'2020-02')
        self.assertIn(('2020-01','CN_UNEMP_SURVEY',5.3),rows)
        self.assertIn(('2020-02','CN_UNEMP_SURVEY',6.2),rows)
        self.assertEqual(len(rows),2)

    def test_first_quarter_explicit_to_range_retains_all_monthly_rates(self):
        html = ('<p>1至3月份，全国城镇调查失业率分别为5.0%、5.0%和5.1%，'
                '分别比上年同月下降0.2、0.4和0.1个百分点；'
                '31个大城市城镇调查失业率分别为4.9%、4.8%和4.9%。</p>')
        self.assertEqual(parse_econ_unemployment_observations(html, '2018-03'), [
            ('2018-01', 'CN_UNEMP_31_CITIES', 4.9),
            ('2018-01', 'CN_UNEMP_SURVEY', 5.0),
            ('2018-02', 'CN_UNEMP_31_CITIES', 4.8),
            ('2018-02', 'CN_UNEMP_SURVEY', 5.0),
            ('2018-03', 'CN_UNEMP_31_CITIES', 4.9),
            ('2018-03', 'CN_UNEMP_SURVEY', 5.1),
        ])

    def test_explicit_to_range_needs_one_rate_for_each_month(self):
        for sentence in (
                '1至3月份，全国城镇调查失业率平均为5.0%。',
                '1至3月份，全国城镇调查失业率分别为5.0%和5.1%。'):
            self.assertEqual(parse_econ_unemployment_observations(
                '<p>' + sentence + '</p>', '2018-03'), [])

    def test_retrospective_monthly_rates_do_not_create_future_observations(self):
        html = '<p>1至3月份，全国城镇调查失业率分别为5.0%、5.0%和5.1%。</p>'
        self.assertEqual(parse_econ_unemployment_observations(html, '2018-02'), [
            ('2018-01', 'CN_UNEMP_SURVEY', 5.0),
            ('2018-02', 'CN_UNEMP_SURVEY', 5.0),
        ])

    def test_incremental_update_stores_explicit_january_in_its_own_period(self):
        html = ('<p>2025年1—2月份主要统计数据。</p>'
                '<p>1月和2月份，全国城镇调查失业率分别为5.1%和5.2%。</p>')
        url = 'https://www.stats.gov.cn/sj/zxfb/202503/verified-fixture.html'
        response = Mock(text=html,content=html.encode('utf-8'),url=url,status_code=200)
        response.iter_content.return_value=[html.encode('utf-8')]
        response.raise_for_status.return_value = None
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / 'update.db')
            with patch('services.china_macro_service.discover_releases', return_value={
                    '2025-02': (url, '2025年1—2月份国民经济起步平稳')}), \
                    patch('services.china_macro_service.sync_annual_macro_snapshots', return_value=0), \
                    patch('services.china_macro_service.requests.get', return_value=response):
                result = update_china_macro(path, start_year=2025, kinds=('econ',), sleep_seconds=0,
                                            archive_root=Path(tmp)/'raw')
            self.assertTrue(result['success'])
            connection = connect(path)
            rows = connection.execute('SELECT period,value,source_url FROM china_macro_observations '
                                      'WHERE indicator_code=? ORDER BY period', ('CN_UNEMP_SURVEY',)).fetchall()
            self.assertEqual([tuple(row) for row in rows], [
                ('2025-01', 5.1, url), ('2025-02', 5.2, url)])
            connection.close()

    def test_incremental_update_does_not_allocate_a_merged_period_average(self):
        html = ('<p>2025年1—2月份主要统计数据。</p>'
                '<p>1—2月份，全国城镇调查失业率平均为5.2%。</p>')
        response = Mock(text=html,content=html.encode('utf-8'),url='https://www.stats.gov.cn/fixture.html',status_code=200)
        response.iter_content.return_value=[html.encode('utf-8')]
        response.raise_for_status.return_value = None
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / 'update.db')
            with patch('services.china_macro_service.discover_releases', return_value={
                    '2025-02': ('https://www.stats.gov.cn/fixture.html', '2025年1—2月份国民经济起步平稳')}), \
                    patch('services.china_macro_service.sync_annual_macro_snapshots', return_value=0), \
                    patch('services.china_macro_service.requests.get', return_value=response):
                result = update_china_macro(path, start_year=2025, kinds=('econ',), sleep_seconds=0,
                                            archive_root=Path(tmp)/'raw')
            self.assertTrue(result['success'])
            connection = connect(path)
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM china_macro_observations '
                                                'WHERE indicator_code=?', ('CN_UNEMP_SURVEY',)).fetchone()[0], 0)
            connection.close()

    def test_subgroup_before_national_sentence_does_not_hide_national_months(self):
        html='<p>一季度，主要就业群体25—59岁失业率均值为4.9%，本地户籍人口失业率均值为5.5%。1、2月份全国城镇调查失业率分别为5.3%、5.5%。</p>'
        rows=parse_econ_unemployment_observations(html,'2022-02')
        self.assertIn(('2022-01','CN_UNEMP_SURVEY',5.3),rows)
        self.assertIn(('2022-02','CN_UNEMP_SURVEY',5.5),rows)
        self.assertNotIn(('2025-01','CN_UNEMP_31_CITIES',5.3),rows)
        self.assertNotIn('2025-01',[row[0] for row in parse_econ_unemployment_observations(
            '1—2月份，全国城镇调查失业率为5.2%。','2025-02')])

    def test_profit_official_times_growth_is_percent_not_two_year_average(self):
        html = ('1—2月份，全国规模以上工业企业实现利润总额11140.1亿元，同比增长1.79倍'
                '（按可比口径计算），比2019年增长72.1%，两年平均增长31.2%。')
        values = parse_industrial_profit_article(html)
        self.assertEqual(values['CN_IND_PROFIT_YTD_YOY'], 179)
        self.assertEqual(values['CN_IND_PROFIT_YTD'], 11140.1)
        annual = parse_industrial_profit_article(
            '全国规模以上工业企业实现利润总额66351.4亿元，比上年增长10.3%。')
        self.assertEqual(annual['CN_IND_PROFIT_YTD_YOY'], 10.3)

    def test_investment_amount_is_not_a_growth_rate(self):
        html = '<table><tr><td>其中：民间投资</td><td>116384</td><td>3.9</td></tr></table>'
        self.assertEqual(parse_fai_article(html)['CN_FAI_PRIVATE_YTD_YOY'], 3.9)
        self.assertNotIn('CN_FAI_PRIVATE_YTD_YOY', parse_fai_article(
            '<table><tr><td>民间投资</td><td>116384</td></tr></table>'))
        self.assertEqual(parse_fai_article('<table><tr><td>民间投资</td><td>-10.1</td></tr></table>')
                         ['CN_FAI_PRIVATE_YTD_YOY'], -10.1)

    def test_cpi_components_use_documented_column_order(self):
        html = '<table><tr><td></td><td>环比涨跌幅（%）</td><td>同比涨跌幅（%）</td></tr>'
        html += '<tr><td>其中：食品</td><td>0.4</td><td>-1.4</td></tr>'
        html += '<tr><td>其中：不包括食品和能源</td><td>0.1</td><td>1.0</td></tr></table>'
        values = parse_cpi_article(html)
        self.assertEqual(values['CN_CPI_FOOD_MOM'], 0.4)
        self.assertEqual(values['CN_CPI_FOOD_YOY'], -1.4)
        self.assertEqual(values['CN_CPI_CORE_YOY'], 1.0)
        self.assertNotIn('CN_CPI_FOOD_YOY', parse_cpi_article('<table><tr><td>食品</td><td>99</td><td>33</td></tr></table>'))

    def test_manufacturing_price_indices_do_not_use_nonmanufacturing_table(self):
        html = '<table><tr><td></td><td>新出口订单</td><td>主要原材料购进价格</td><td>出厂价格</td></tr>'
        html += '<tr><td>2026年8月</td><td>49.2</td><td>60.5</td><td>55.8</td></tr>'
        html += '<tr><td>2026年9月</td><td>50.0</td><td>60.8</td><td>54.0</td></tr></table>'
        html += '<table><tr><td></td><td>新出口订单</td><td>投入品价格</td><td>销售价格</td></tr>'
        html += '<tr><td>2026年9月</td><td>48.5</td><td>52.5</td><td>50.3</td></tr></table>'
        values = parse_pmi_article(html)
        self.assertEqual(values['CN_PMI_MFG_NEW_EXPORT_ORDERS'], 50.0)
        self.assertEqual(values['CN_PMI_MFG_INPUT_PRICES'], 60.8)

    def test_profit_tables_keep_precise_values_and_official_yoy(self):
        html = '<table><tr><td>分组</td><td>营业收入</td><td>营业成本</td><td>利润总额</td></tr>'
        html += '<tr><td>总计</td><td>930851.2</td><td>6.6</td><td>791893.6</td><td>6.1</td><td>52719.8</td><td>15.7</td></tr></table>'
        html += '<p>规模以上工业企业实现营业收入93.09万亿元，同比增长6.6%。</p>'
        html += '<p>8月末，规模以上工业企业应收账款29.48万亿元，同比增长9.0%；产成品存货7.36万亿元，增长11.0%。</p>'
        values = parse_industrial_profit_article(html)
        self.assertEqual(values['CN_IND_REVENUE_YTD'], 930851.2)
        self.assertEqual(values['CN_IND_PROFIT_YTD_YOY'], 15.7)
        self.assertEqual(values['CN_IND_RECEIVABLES'], 294800)
        self.assertEqual(values['CN_IND_FINISHED_INVENTORY_YOY'], 11.0)

    def test_annual_age_definitions_and_real_income_are_not_substituted(self):
        html = '<p>死亡人口1093万人。全年城镇新增就业1256万人。16—59岁人口85798万人。</p>'
        html += '<p>全国居民人均可支配收入41314元，比上年名义增长5.3%，扣除价格因素实际增长5.1%。</p>'
        html += '<p>全国居民人均消费支出28227元，比上年名义增长5.3%，扣除价格因素实际增长5.1%。</p>'
        values = parse_annual_livelihood_article(html)
        self.assertEqual(values['CN_DEATHS'], 1093)
        self.assertEqual(values['CN_WORKING_AGE_POP'], 85798)
        self.assertEqual(values['CN_INCOME_REAL_YOY'], 5.1)
        self.assertEqual(values['CN_CONSUMPTION_REAL_YOY'], 5.1)
        self.assertNotIn('CN_WORKING_AGE_POP', parse_annual_livelihood_article('15—59岁人口89438万人。'))

    def test_merged_january_is_not_a_gap_but_missing_march_is(self):
        merged = calendar_coverage(['2024-02','2024-04','2025-02'], 'monthly', True)
        self.assertIn('2024-03', merged['missing_periods'])
        self.assertNotIn('2025-01', merged['missing_periods'])
        self.assertIsNone(calendar_coverage(['2026-09-25','2026-09-28'], 'daily')['internal_gap_count'])

    def test_us_changes_require_calendar_comparison_periods(self):
        rows = [{'period':'2025-01-01','value':100}, {'period':'2026-01-01','value':103}]
        self.assertEqual(derive_change_series(rows,12)[0]['value'],3)
        self.assertEqual(derive_difference_series(rows),[])
        missing_quarter = [{'period':'2025-07-01','value':100}, {'period':'2026-01-01','value':103}]
        self.assertEqual(derive_change_series(missing_quarter,1,True),[])

    def test_new_us_series_and_release_frequency(self):
        registry = {row[0]:row for row in SERIES}
        for code in ('PCEPI','DSPIC96','PCEC96','PSAVERT','CCSA','TOTBKCR','HOUST','PERMIT','HSN1F','MSPUS','EXPGS','IMPGS'):
            self.assertIn(code,registry)
        self.assertEqual(registry['MSPUS'][3],'quarterly')
        self.assertEqual(registry['TOTBKCR'][3],'weekly')

    def test_balance_sheet_empty_cells_keep_their_actual_columns(self):
        wb=Workbook(); ws=wb.active
        ws['A1']='货币当局资产负债表'; ws['C2']='2026.08'; ws['E2']='2026.09'
        ws['A4']='总资产';ws['C4']=450000;ws['E4']=460000
        stream=io.BytesIO();wb.save(stream)
        rows=parse_balance_sheet_xlsx(stream.getvalue(),{'source_url':'https://example.test/balance.xlsx'})
        self.assertEqual({row['period']:row['value'] for row in rows if row['indicator_code']=='total_assets'},
                         {'2026-08':450000,'2026-09':460000})

    def test_balance_sheet_vintages_never_sum(self):
        rows=[{'indicator_code':'total_assets','period':'2026-05','value':400000,'published_date':'2026-06-01'},
              {'indicator_code':'total_assets','period':'2026-05','value':410000,'published_date':'2026-07-01'}]
        self.assertEqual(len(latest_balance_vintages(rows)),1)
        self.assertEqual(latest_balance_vintages(rows)[0]['value'],410000)

    def test_empty_quality_is_missing_not_zero_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            result=build_macro_data_quality(str(Path(tmp)/'empty.db'),'china')
        self.assertEqual(result['series_with_data'],0)
        self.assertIn('CN_IND_PROFIT_YTD',result['missing_series'])

    def test_release_periods_preserve_merged_months_and_quarters(self):
        self.assertEqual(profit_period('2026年1—2月份全国规模以上工业企业利润增长8%'),'2026-02')
        self.assertEqual(capacity_period('2026年二季度全国工业产能利用率为75%'),'2026-06')

    def test_migration_address_is_not_a_reporting_year(self):
        url='https://www.stats.gov.cn/sj/zxfb/202302/t20230203_1900718.html'
        self.assertIsNone(_econ_period('4月份国民经济运行继续改善',url))
        self.assertEqual(econ_period_from_article('附表2020年1-4月份主要指标数据',
                                                 '4月份国民经济运行继续改善',url),'2020-04')

    def test_annual_footnote_markers_do_not_hide_median_or_migrants(self):
        values=parse_annual_livelihood_article('全国居民人均可支配收入中位数<sup>[5]</sup>35910元。全国农民工<sup>[12]</sup>总量30115万人。')
        self.assertEqual(values['CN_DISPOSABLE_INCOME_MEDIAN'],35910)
        self.assertEqual(values['CN_MIGRANT_WORKERS'],30115)

    def test_reviewed_bank_table_has_no_future_or_synthetic_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            conn=connect(str(Path(tmp)/'bank.db'));ensure_china_macro_tables(conn)
            history=reviewed_history_rows(HISTORY_PATH)
            self.assertEqual(sync_bank_supervision_snapshot(conn)['records_changed'],28+len(history))
            self.assertEqual(sync_bank_supervision_snapshot(conn)['records_changed'],0)
            periods={row[0] for row in conn.execute("SELECT DISTINCT period FROM china_macro_observations")}
            self.assertEqual(periods,{'2026-03','2026-06'}|{row['period'] for row in history})
            self.assertNotIn('2026-09',periods)
            value=conn.execute("SELECT value FROM china_macro_observations WHERE indicator_code='CN_BANK_NIM_YTD' AND period='2026-06'").fetchone()[0]
            self.assertEqual(value,1.41)
            row=conn.execute("SELECT value,release_date,source_url FROM china_macro_observations WHERE indicator_code='CN_BANK_NET_PROFIT_YTD' AND period='2025-12'").fetchone()
            self.assertEqual(row[0],23776)
            self.assertEqual(row[1],'2026-02-12')
            self.assertIn('1208453',row[2])
            conn.close()

    def test_custom_bank_snapshot_does_not_mix_production_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'snapshot.csv'
            path.write_bytes(SNAPSHOT_PATH.read_bytes())
            conn=connect(str(Path(tmp)/'bank.db'));ensure_china_macro_tables(conn)
            self.assertEqual(sync_bank_supervision_snapshot(conn,path=path)['rows_reviewed'],28)
            conn.close()

    def test_reviewed_bank_history_keeps_blank_quarters_missing(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'review.json'
            table={'year':2024,'reviewed_at':'2026-10-10','source_url':'https://www.nfra.gov.cn/table.png',
                   'source_title':'Official table','release_date':'2025-02-21','retrieved_at':'2026-10-10',
                   'parser_notes':'Reviewed quarter cells','values':{'CN_BANK_NPL_RATIO':[1.59,None,1.56,1.50]}}
            path.write_text(json.dumps({'schema_version':1,'tables':[table]}),encoding='utf-8')
            rows=reviewed_history_rows(path)
            self.assertEqual({row['period'] for row in rows},{'2024-03','2024-09','2024-12'})

    def test_econ_cumulative_core_cpi_cannot_overwrite_monthly_core(self):
        html='1—7月份，扣除食品和能源价格后的核心CPI同比上涨1.1%。其中，7月份核心CPI同比上涨0.9%。'
        values=parse_econ_article(html,period='2026-07')
        self.assertNotIn('CN_CPI_CORE_YOY',values)
        self.assertEqual(values['CN_CPI_CORE_YTD_AVG_YOY'],1.1)
        jan='附表2026年1-2月份主要指标数据。1月份，全国城镇调查失业率为5.2%。'
        self.assertNotIn('CN_UNEMP_SURVEY',parse_econ_article(jan,period='2026-02'))


if __name__ == '__main__':
    unittest.main()
